"""Accounting reconciliation and separate-format CMS cross-table validation."""

from __future__ import annotations

import io
import zipfile
from collections import defaultdict
from typing import Any, Iterable

import openpyxl
import pandas as pd

from reconciliation_healthcare.normalize_nhea import TABLES_ARCHIVE
from reconciliation_healthcare.paths import (
    LEDGER_CELLS_PARQUET_PATH as LEDGER_CELLS_PATH,
    LEDGER_PARQUET_PATH,
    RAW_DIR,
    REPORT_PATH,
    SOURCE_SERVICE_PATH,
    SPONSOR_PATH,
)


MILLION = 1_000_000

PHC_SERVICE_CODES = (
    "hospital_care",
    "physician_clinical_services",
    "other_professional_services",
    "dental_services",
    "other_health_residential_personal_care",
    "home_health_care",
    "nursing_care_ccrc",
    "prescription_drugs",
    "other_nondurable_medical_products",
    "durable_medical_equipment",
)
DISTRIBUTED_SERVICE_CODES = PHC_SERVICE_CODES + ("administration_nonmedical_insurance",)
SOURCE_HIERARCHY_SERVICE_CODES = (
    "nhe",
    "hce",
    "phc",
    "hospital_care",
    "physician_clinical_services",
    "dental_services",
    "other_professional_services",
    "home_health_care",
    "other_nondurable_medical_products",
    "prescription_drugs",
    "durable_medical_equipment",
    "nursing_care_ccrc",
    "other_health_residential_personal_care",
    "administration_nonmedical_insurance",
    "state_local_administration",
    "federal_administration",
    "nonmedical_insurance",
)
ADDITIVE_FUNDING_CODES = (
    "out_of_pocket",
    "private_health_insurance",
    "medicare",
    "medicaid",
    "chip",
    "department_of_defense",
    "department_of_veterans_affairs",
    "other_third_party_payers_programs",
)
NHE_SOURCE_CODES = ADDITIVE_FUNDING_CODES + (
    "government_public_health_activities",
    "investment",
)

# Each tuple is (reported parent, mutually exclusive children, rounding tolerance).
# The tolerances reflect sums of independently rounded integer-million source cells.
SOURCE_HIERARCHIES = (
    (
        "health_insurance",
        (
            "private_health_insurance",
            "medicare",
            "medicaid",
            "chip",
            "department_of_defense",
            "department_of_veterans_affairs",
        ),
        3 * MILLION,
    ),
    ("medicaid", ("medicaid_federal", "medicaid_state_local"), MILLION),
    ("chip", ("chip_federal", "chip_state_local"), MILLION),
    (
        "maternal_child_health",
        ("maternal_child_health_federal", "maternal_child_health_state_local"),
        MILLION,
    ),
    (
        "vocational_rehabilitation",
        (
            "vocational_rehabilitation_federal",
            "vocational_rehabilitation_state_local",
        ),
        MILLION,
    ),
    (
        "other_third_party_payers_programs",
        (
            "worksite_health_care",
            "other_private_revenues",
            "indian_health_service",
            "workers_compensation",
            "general_assistance",
            "maternal_child_health",
            "vocational_rehabilitation",
            "other_federal_programs",
            "samhsa",
            "other_state_local_programs",
            "school_health",
        ),
        6 * MILLION,
    ),
    ("total_cms_programs", ("medicare", "medicaid", "chip"), 2 * MILLION),
)

KNOWN_2024_USD = {
    ("nhe", "all_sources"): 5_278_588 * MILLION,
    ("nhe", "medicare"): 1_118_000 * MILLION,
    ("nhe", "medicaid"): 931_692 * MILLION,
    ("nhe", "private_health_insurance"): 1_644_592 * MILLION,
    ("nhe", "out_of_pocket"): 556_554 * MILLION,
    ("hospital_care", "all_sources"): 1_634_738 * MILLION,
    ("physician_clinical_services", "all_sources"): 1_109_681 * MILLION,
    ("prescription_drugs", "all_sources"): 466_968 * MILLION,
}

EXPECTED_ACCOUNTING_CHECK_COUNT = 65 * (6 + len(DISTRIBUTED_SERVICE_CODES))
EXPECTED_SOURCE_HIERARCHY_CHECK_COUNT = (
    65 * len(SOURCE_HIERARCHY_SERVICE_CODES) * len(SOURCE_HIERARCHIES)
)
EXPECTED_VALIDATION_CHECK_COUNT = (
    EXPECTED_ACCOUNTING_CHECK_COUNT
    + EXPECTED_SOURCE_HIERARCHY_CHECK_COUNT
    + 38  # sponsor years, 1987-2024
    + 1  # generated ledger cross-dimensional total
    + 110  # separate-format Table 19 cells and totals
    + len(KNOWN_2024_USD)
)


def _amount(
    dataframe: pd.DataFrame,
    year: int,
    service_code: str,
    funding_code: str,
    *,
    null_as_zero: bool = False,
) -> int | None:
    matches = dataframe[
        (dataframe["year"] == year)
        & (dataframe["service_category_code"] == service_code)
        & (dataframe["funding_source_code"] == funding_code)
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one normalized record for {year}/{service_code}/{funding_code}; "
            f"found {len(matches)}"
        )
    value = matches.iloc[0]["amount_usd"]
    if pd.isna(value):
        return 0 if null_as_zero else None
    return int(value)


def _sum_amounts(values: Iterable[int | None]) -> int:
    return sum(0 if value is None else value for value in values)


def check(
    identity: str,
    year: int,
    reported_usd: int,
    calculated_usd: int,
    tolerance_usd: int,
    *,
    source: str = "NHE2024.csv",
) -> dict[str, Any]:
    residual = calculated_usd - reported_usd
    return {
        "identity": identity,
        "year": year,
        "reported_usd": reported_usd,
        "calculated_usd": calculated_usd,
        "residual_usd": residual,
        "tolerance_usd": tolerance_usd,
        "passed": abs(residual) <= tolerance_usd,
        "validation_source": source,
    }


def accounting_checks(source_service: pd.DataFrame) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    for year in range(1960, 2025):
        nhe = _amount(source_service, year, "nhe", "all_sources")
        hce = _amount(source_service, year, "hce", "all_sources")
        phc = _amount(source_service, year, "phc", "all_sources")
        admin = _amount(
            source_service, year, "administration_nonmedical_insurance", "all_sources"
        )
        public_health = _amount(source_service, year, "public_health", "all_sources")
        investment = _amount(source_service, year, "nhe", "investment")
        assert all(value is not None for value in (nhe, hce, phc, admin, public_health, investment))

        phc_services = _sum_amounts(
            _amount(source_service, year, service, "all_sources")
            for service in PHC_SERVICE_CODES
        )
        nhe_services = phc_services + int(admin) + int(public_health) + int(investment)
        results.append(check("NHE from mutually exclusive service components", year, int(nhe), nhe_services, 10 * MILLION))

        nhe_sources = _sum_amounts(
            _amount(source_service, year, "nhe", source, null_as_zero=True)
            for source in NHE_SOURCE_CODES
        )
        results.append(check("NHE from mutually exclusive source components", year, int(nhe), nhe_sources, 5 * MILLION))
        results.append(check("PHC from ten service components", year, int(phc), phc_services, 5 * MILLION))
        results.append(
            check(
                "HCE from PHC, administration/non-medical insurance, and public health",
                year,
                int(hce),
                int(phc) + int(admin) + int(public_health),
                3 * MILLION,
            )
        )
        results.append(
            check(
                "Investment from research and structures/equipment",
                year,
                int(investment),
                _sum_amounts(
                    (
                        _amount(source_service, year, "research", "all_sources"),
                        _amount(source_service, year, "structures_equipment", "all_sources"),
                    )
                ),
                2 * MILLION,
            )
        )
        results.append(
            check(
                "Administration/non-medical insurance from three components",
                year,
                int(admin),
                _sum_amounts(
                    _amount(source_service, year, service, "all_sources")
                    for service in (
                        "state_local_administration",
                        "federal_administration",
                        "nonmedical_insurance",
                    )
                ),
                3 * MILLION,
            )
        )

        for service in DISTRIBUTED_SERVICE_CODES:
            reported = _amount(source_service, year, service, "all_sources")
            assert reported is not None
            calculated = _sum_amounts(
                _amount(source_service, year, service, source, null_as_zero=True)
                for source in ADDITIVE_FUNDING_CODES
            )
            results.append(
                check(
                    f"Service funding distribution: {service}",
                    year,
                    int(reported),
                    calculated,
                    5 * MILLION,
                )
            )
    return results


def source_hierarchy_checks(source_service: pd.DataFrame) -> list[dict[str, Any]]:
    """Reconcile every reviewed payer subtree in every full service panel/year."""
    key_columns = ["year", "service_category_code", "funding_source_code"]
    if source_service.duplicated(key_columns).any():
        raise ValueError("Source/service data contains duplicate canonical keys")
    amount_lookup = {
        (int(row.year), str(row.service_category_code), str(row.funding_source_code)): (
            0 if pd.isna(row.amount_usd) else int(row.amount_usd)
        )
        for row in source_service.itertuples(index=False)
    }

    def amount(year: int, service_code: str, funding_code: str) -> int:
        key = (year, service_code, funding_code)
        try:
            return amount_lookup[key]
        except KeyError as error:
            raise ValueError(
                f"Missing normalized record for {year}/{service_code}/{funding_code}"
            ) from error

    results: list[dict[str, Any]] = []
    for year in range(1960, 2025):
        for service_code in SOURCE_HIERARCHY_SERVICE_CODES:
            for parent_code, child_codes, tolerance_usd in SOURCE_HIERARCHIES:
                reported = amount(year, service_code, parent_code)
                calculated = sum(
                    amount(year, service_code, child_code) for child_code in child_codes
                )
                result = check(
                    f"Source hierarchy: {service_code}/{parent_code}",
                    year,
                    reported,
                    calculated,
                    tolerance_usd,
                )
                result["service_category_code"] = service_code
                result["funding_source_code"] = parent_code
                results.append(result)

    if len(results) != EXPECTED_SOURCE_HIERARCHY_CHECK_COUNT:
        raise AssertionError(
            "Detailed source hierarchy validation produced "
            f"{len(results):,} checks; expected {EXPECTED_SOURCE_HIERARCHY_CHECK_COUNT:,}"
        )
    return results


def sponsor_checks(sponsor: pd.DataFrame) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    for year in range(1987, 2025):
        rows = sponsor[sponsor["year"] == year]
        reported = int(rows.loc[rows["sponsor_code"] == "all_sponsors", "amount_usd"].iloc[0])
        calculated = int(rows.loc[rows["include_in_sponsor_sum"], "amount_usd"].sum())
        results.append(
            check(
                "Sponsor total from five mutually exclusive sponsors",
                year,
                reported,
                calculated,
                400_000_000,
                source="NHE Tables, Table 5 (rounded to $0.1B)",
            )
        )
    return results


def ledger_cross_check(source_service: pd.DataFrame) -> list[dict[str, Any]]:
    missing = [path for path in (LEDGER_CELLS_PATH, LEDGER_PARQUET_PATH) if not path.exists()]
    if missing:
        raise FileNotFoundError(f"Required generated ledger artifact is missing: {missing[0]}")

    # Compare the persisted artifacts to a fresh programmatic rendering. This
    # makes the validation sensitive to a stale or manually altered cell while
    # retaining one numerical cross-dimensional identity in the result count.
    from reconciliation_healthcare.ledger import build_ledger_cells, build_wide_ledger

    expected_cells = build_ledger_cells(source_service)
    cells = pd.read_parquet(LEDGER_CELLS_PATH)
    expected_wide = build_wide_ledger(expected_cells)
    wide = pd.read_parquet(LEDGER_PARQUET_PATH)
    try:
        pd.testing.assert_frame_equal(cells, expected_cells, check_dtype=True)
        pd.testing.assert_frame_equal(wide, expected_wide, check_dtype=True)
    except AssertionError as error:
        raise ValueError(
            "Persisted 2024 ledger artifacts do not exactly match programmatic regeneration"
        ) from error

    interior = cells[
        (cells["ledger_row_code"] != "total")
        & (cells["ledger_column_code"] != "total")
    ]
    calculated = int(interior["amount_usd"].sum())
    reported = int(_amount(source_service, 2024, "nhe", "all_sources") or 0)
    return [
        check(
            "2024 ledger interior cross-dimensional total",
            2024,
            reported,
            calculated,
            10 * MILLION,
            source="Generated ledger from NHE2024.csv",
        )
    ]


def _table19_value(worksheet: Any, row: int, columns: int | tuple[int, ...]) -> int:
    column_numbers = (columns,) if isinstance(columns, int) else columns
    total = 0
    for column in column_numbers:
        value = worksheet.cell(row, column).value
        if value is None or (isinstance(value, str) and value.strip() in {"", "—", "-", "–"}):
            continue
        if not isinstance(value, (int, float)):
            raise ValueError(f"Unexpected Table 19 value at row {row}, column {column}: {value!r}")
        total += int(value) * MILLION
    return total


def table19_checks(source_service: pd.DataFrame) -> list[dict[str, Any]]:
    member = "Table 19 National Health Expenditures by Type of Expenditure and Program.xlsx"
    with zipfile.ZipFile(RAW_DIR / TABLES_ARCHIVE) as archive:
        workbook = openpyxl.load_workbook(
            io.BytesIO(archive.read(member)), read_only=True, data_only=True
        )
    if workbook.sheetnames != ["Table 19"]:
        raise ValueError(f"Unexpected Table 19 sheets: {workbook.sheetnames}")
    worksheet = workbook["Table 19"]
    service_columns: dict[str, int | tuple[int, ...]] = {
        "hospital_care": 8,
        "physician_clinical_services": 9,
        "other_professional_services": 10,
        "dental_services": 11,
        "other_health_residential_personal_care": 13,
        "home_health_care": 14,
        "nursing_care_ccrc": 15,
        "prescription_drugs": 17,
        "durable_medical_equipment": 18,
        "other_nondurable_medical_products": 19,
        "administration_nonmedical_insurance": (22, 23, 24),
    }
    source_rows = {
        "out_of_pocket": 9,
        "private_health_insurance": 11,
        "medicare": 12,
        "medicaid": 13,
        "chip": 16,
        "department_of_defense": 19,
        "department_of_veterans_affairs": 20,
        "other_third_party_payers_programs": 21,
    }
    results: list[dict[str, Any]] = []
    source_name = "Separate-format CMS NHE Tables, Table 19 cross-check"
    for service_code, columns in service_columns.items():
        for funding_code, row in source_rows.items():
            reported = _table19_value(worksheet, row, columns)
            calculated = int(
                _amount(
                    source_service,
                    2024,
                    service_code,
                    funding_code,
                    null_as_zero=True,
                )
                or 0
            )
            results.append(
                check(
                    f"Table 19 cell: {service_code}/{funding_code}",
                    2024,
                    reported,
                    calculated,
                    MILLION if service_code == "administration_nonmedical_insurance" else 0,
                    source=source_name,
                )
            )

    for service_code, columns in service_columns.items():
        reported = _table19_value(worksheet, 8, columns)
        calculated = int(_amount(source_service, 2024, service_code, "all_sources") or 0)
        results.append(
            check(
                f"Table 19 service total: {service_code}",
                2024,
                reported,
                calculated,
                0,
                source=source_name,
            )
        )
    for funding_code, row in source_rows.items():
        reported = _table19_value(worksheet, row, 7)
        calculated = int(_amount(source_service, 2024, "nhe", funding_code) or 0)
        results.append(
            check(
                f"Table 19 source total: {funding_code}",
                2024,
                reported,
                calculated,
                0,
                source=source_name,
            )
        )
    for identity, row, column, service_code, funding_code in (
        ("Table 19 NHE total", 8, 7, "nhe", "all_sources"),
        ("Table 19 public-health total", 37, 25, "public_health", "all_sources"),
        ("Table 19 investment total", 40, 7, "nhe", "investment"),
    ):
        results.append(
            check(
                identity,
                2024,
                _table19_value(worksheet, row, column),
                int(_amount(source_service, 2024, service_code, funding_code) or 0),
                0,
                source=source_name,
            )
        )
    return results


def headline_checks(source_service: pd.DataFrame) -> list[dict[str, Any]]:
    results = []
    for (service_code, funding_code), expected in KNOWN_2024_USD.items():
        actual = int(_amount(source_service, 2024, service_code, funding_code) or 0)
        results.append(
            check(
                f"2024 published sanity value: {service_code}/{funding_code}",
                2024,
                expected,
                actual,
                MILLION,
                source="CMS published headline checked against principal download",
            )
        )
    return results


def run_validation() -> list[dict[str, Any]]:
    source_service = pd.read_parquet(SOURCE_SERVICE_PATH)
    sponsor = pd.read_parquet(SPONSOR_PATH)
    results = [
        *accounting_checks(source_service),
        *source_hierarchy_checks(source_service),
        *sponsor_checks(sponsor),
        *ledger_cross_check(source_service),
        *table19_checks(source_service),
        *headline_checks(source_service),
    ]
    if len(results) != EXPECTED_VALIDATION_CHECK_COUNT:
        raise AssertionError(
            f"Validation produced {len(results):,} checks; "
            f"expected the full set of {EXPECTED_VALIDATION_CHECK_COUNT:,}"
        )
    return results


def _billions(value: int) -> str:
    return f"${value / 1_000_000_000:,.3f}B"


def write_reconciliation_report(results: list[dict[str, Any]]) -> None:
    source_service = pd.read_parquet(SOURCE_SERVICE_PATH)
    sponsor = pd.read_parquet(SPONSOR_PATH)

    def amount(service: str, source: str) -> int:
        return int(_amount(source_service, 2024, service, source) or 0)

    source_rows = [
        ("Out of pocket", amount("nhe", "out_of_pocket")),
        ("Private Health Insurance", amount("nhe", "private_health_insurance")),
        ("Medicare", amount("nhe", "medicare")),
        ("Medicaid", amount("nhe", "medicaid")),
        ("CHIP", amount("nhe", "chip")),
        ("Department of Defense", amount("nhe", "department_of_defense")),
        ("Department of Veterans Affairs", amount("nhe", "department_of_veterans_affairs")),
        (
            "Other Third Party Payers and Programs",
            amount("nhe", "other_third_party_payers_programs"),
        ),
        (
            "Government Public Health Activities",
            amount("nhe", "government_public_health_activities"),
        ),
        ("Investment", amount("nhe", "investment")),
    ]
    service_rows = [
        (source_service[
            (source_service["year"] == 2024)
            & (source_service["service_category_code"] == service)
            & (source_service["funding_source_code"] == "all_sources")
        ].iloc[0]["service_category_name"], amount(service, "all_sources"))
        for service in PHC_SERVICE_CODES + ("administration_nonmedical_insurance", "public_health")
    ] + [("Investment", amount("nhe", "investment"))]

    by_identity: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for result in results:
        by_identity[result["identity"]].append(result)
    core_identities = [
        "NHE from mutually exclusive service components",
        "NHE from mutually exclusive source components",
        "PHC from ten service components",
        "HCE from PHC, administration/non-medical insurance, and public health",
        "Investment from research and structures/equipment",
        "Administration/non-medical insurance from three components",
        "Sponsor total from five mutually exclusive sponsors",
        "2024 ledger interior cross-dimensional total",
    ]

    sponsor_2024 = sponsor[sponsor["year"] == 2024]
    sponsor_leaves = sponsor_2024[sponsor_2024["include_in_sponsor_sum"]]
    table19 = [result for result in results if result["identity"].startswith("Table 19")]
    source_hierarchy = [
        result for result in results if result["identity"].startswith("Source hierarchy:")
    ]
    failed = [result for result in results if not result["passed"]]
    nhe = amount("nhe", "all_sources")

    lines = [
        "# Reconciliation report",
        "",
        "## 2024 national total",
        "",
        f"The principal CMS file reports **{_billions(nhe)}** "
        f"(${nhe:,}) of National Health Expenditures for calendar year 2024.",
        "",
        "## Totals by source/component",
        "",
        "| Source/component | Current USD (exact conversion of reported millions) |",
        "|---|---:|",
        *[f"| {name} | ${value:,} |" for name, value in source_rows],
        f"| **Total NHE** | **${nhe:,}** |",
        "",
        "The source rows above are mutually exclusive at this ledger level. Health Insurance and "
        "its components, and Other Third Party Payers and Programs and its children, are retained "
        "in normalized data but are not summed together. Government Public Health Activities and "
        "Investment are expenditure components used by CMS to complete the full-NHE matrix.",
        "",
        "## Totals by top-level service",
        "",
        "| Service | Current USD (exact conversion of reported millions) |",
        "|---|---:|",
        *[f"| {name} | ${value:,} |" for name, value in service_rows],
        f"| **Total NHE** | **${nhe:,}** |",
        "",
        "## Accounting identities",
        "",
        "| Identity | 2024 residual | Historical maximum absolute residual | Tolerance | Status |",
        "|---|---:|---:|---:|---|",
    ]
    for identity in core_identities:
        identity_results = by_identity.get(identity, [])
        if not identity_results:
            continue
        result_2024 = next(
            (result for result in identity_results if result["year"] == 2024),
            identity_results[-1],
        )
        maximum = max(abs(result["residual_usd"]) for result in identity_results)
        status = "PASS" if all(result["passed"] for result in identity_results) else "FAIL"
        lines.append(
            f"| {identity} | ${result_2024['residual_usd']:,} | ${maximum:,} | "
            f"${result_2024['tolerance_usd']:,} | {status} |"
        )

    lines.extend(
        [
            "",
            f"The separate-format exact-million CMS Table 19 check compared **{len(table19)}** 2024 "
            f"cells/totals; **{sum(result['passed'] for result in table19)} passed**. Direct cells use "
            "zero tolerance; combined administration cells allow $1 million because they sum three "
            "separately rounded Table 19 components.",
            "",
            "Before the ledger identity is evaluated, validation requires both persisted 2024 "
            "ledger Parquets to match a fresh programmatic regeneration exactly, including every "
            "cell-level provenance field.",
            "",
            f"The historical source hierarchy suite compared **{len(source_hierarchy):,}** "
            "parent/child totals across seven payer subtrees, 17 service panels, and 65 years; "
            f"**{sum(result['passed'] for result in source_hierarchy):,} passed**. Its tolerances "
            "range from $1 million to $6 million based on the number of independently rounded "
            "integer-million children.",
            "",
            "## Sponsor view (separate accounting perspective)",
            "",
            "| Sponsor | CMS-reported current USD |",
            "|---|---:|",
            *[
                f"| {row.sponsor_name} | ${int(row.amount_usd):,} |"
                for row in sponsor_leaves.itertuples(index=False)
            ],
            "",
            "Table 5 reports sponsor amounts to $0.1 billion. Its five 2024 leaf sponsors total "
            f"${int(sponsor_leaves['amount_usd'].sum()):,}, versus the rounded sponsor root of "
            f"${int(sponsor_2024.loc[sponsor_2024['sponsor_code'] == 'all_sponsors', 'amount_usd'].iloc[0]):,}. "
            "The $0.1 billion residual is published rounding, not an allocated adjustment.",
            "",
            "Sponsor and source of funds are alternative perspectives on overlapping expenditures; "
            "they are not added or synthetically crossed by service.",
            "",
            "## Rounding and source-table differences",
            "",
            "- The principal historical CSV reports integer millions in current dollars. Summing many "
            "reported cells can therefore produce residuals of a few million dollars.",
            "- Table 4 reports billions and is used for presentation review, not exact reconciliation. "
            "Table 19 reports 2024 levels in millions and is a separate-format cross-check from the "
            "same CMS release, not an independent estimate.",
            "- The NHE Summary reports amounts in billions, population in millions, per-capita dollars, "
            "growth/distribution percentages, and GDP share. Those non-dollar rows are not ingested as flows.",
            "- An explicit source dash is retained as `not_applicable`; an empty source cell is retained "
            "as `structural_blank`. Both have null amounts and are never converted to reported zeroes. "
            "Blank wide-ledger crossings are likewise null.",
            "",
            "## Open accounting questions",
            "",
            "- The complete service/source matrix necessarily uses Government Public Health Activities "
            "and Investment as top-level NHE components, not payer programs comparable to Medicare.",
            "- Table 5 sponsor history begins in 1987; CMS does not supply a 1960–2024 service × sponsor cube.",
            "- The 2024 comprehensive benchmark revision reopens the full history. A later vintage should "
            "be ingested as a new pinned release rather than silently replacing these checksums.",
            "- Provider-establishment classification means `service_category` is not necessarily the "
            "clinical procedure performed or the legal recipient of each payment.",
            "",
            "## Validation result",
            "",
            f"**{'PASS' if not failed else 'FAIL'}:** {len(results) - len(failed):,} of "
            f"{len(results):,} checks passed.",
            "",
        ]
    )
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text("\n".join(lines), encoding="utf-8")


def validate_and_write_report() -> list[dict[str, Any]]:
    results = run_validation()
    write_reconciliation_report(results)
    return results


def main() -> None:
    results = validate_and_write_report()
    failed = [result for result in results if not result["passed"]]
    print(f"{len(results) - len(failed):,}/{len(results):,} reconciliation checks passed")
    print(f"Wrote {REPORT_PATH}")
    if failed:
        for result in failed[:20]:
            print(
                f"FAIL {result['identity']} ({result['year']}): "
                f"residual={result['residual_usd']:,}, tolerance={result['tolerance_usd']:,}"
            )
        raise SystemExit(1)


if __name__ == "__main__":
    main()
