"""Normalize the inspected CMS NHEA accounting views with cell-level provenance."""

from __future__ import annotations

import csv
import io
import json
import zipfile
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any

import duckdb
import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

from reconciliation_healthcare.paths import (
    DUCKDB_PATH,
    MANIFEST_PATH,
    RAW_DIR,
    SOURCE_SERVICE_PATH,
    SPONSOR_PATH,
)


PRINCIPAL_ARCHIVE = "national-health-expenditures-type-service-source-funds-cy-1960-2024.zip"
PRINCIPAL_MEMBER = "NHE2024.csv"
TABLES_ARCHIVE = "nhe-tables.zip"
SPONSOR_MEMBER = "Table 05 National Health Expenditures by Type of Sponsor.xlsx"
@dataclass(frozen=True)
class ServiceMeta:
    code: str
    name: str
    parent: str | None
    level: int
    is_subtotal: bool
    include_in_nhe_service_sum: bool
    ledger_order: int | None


@dataclass(frozen=True)
class FundingMeta:
    code: str
    name: str
    parent: str | None
    level: int
    is_subtotal: bool
    include_in_service_funding_sum: bool
    ledger_order: int | None


@dataclass(frozen=True)
class SponsorMeta:
    code: str
    name: str
    parent: str | None
    level: int
    is_subtotal: bool
    include_in_sponsor_sum: bool


SERVICES: dict[str, ServiceMeta] = {
    item.code: item
    for item in (
        ServiceMeta("nhe", "National Health Expenditures", None, 0, True, False, None),
        ServiceMeta("hce", "Health Consumption Expenditures", "nhe", 1, True, False, None),
        ServiceMeta("phc", "Personal Health Care", "hce", 2, True, False, None),
        ServiceMeta("hospital_care", "Hospital Care", "phc", 3, False, True, 10),
        ServiceMeta(
            "physician_clinical_services",
            "Physician and Clinical Services",
            "phc",
            3,
            False,
            True,
            20,
        ),
        ServiceMeta("other_professional_services", "Other Professional Services", "phc", 3, False, True, 30),
        ServiceMeta("dental_services", "Dental Services", "phc", 3, False, True, 40),
        ServiceMeta(
            "other_health_residential_personal_care",
            "Other Health, Residential, and Personal Care",
            "phc",
            3,
            False,
            True,
            50,
        ),
        ServiceMeta("home_health_care", "Home Health Care", "phc", 3, False, True, 60),
        ServiceMeta(
            "nursing_care_ccrc",
            "Nursing Care Facilities and Continuing Care Retirement Communities",
            "phc",
            3,
            False,
            True,
            70,
        ),
        ServiceMeta("prescription_drugs", "Retail Prescription Drugs", "phc", 3, False, True, 80),
        ServiceMeta(
            "other_nondurable_medical_products",
            "Other Non-Durable Medical Products",
            "phc",
            3,
            False,
            True,
            90,
        ),
        ServiceMeta("durable_medical_equipment", "Durable Medical Equipment", "phc", 3, False, True, 100),
        ServiceMeta(
            "administration_nonmedical_insurance",
            "Government Administration and Non-Medical Insurance Expenditures",
            "hce",
            2,
            True,
            True,
            110,
        ),
        ServiceMeta(
            "state_local_administration",
            "State and Local Government Administration",
            "administration_nonmedical_insurance",
            3,
            False,
            False,
            None,
        ),
        ServiceMeta(
            "federal_administration",
            "Federal Government Administration",
            "administration_nonmedical_insurance",
            3,
            False,
            False,
            None,
        ),
        ServiceMeta(
            "nonmedical_insurance",
            "Non-Medical Insurance Expenditures",
            "administration_nonmedical_insurance",
            3,
            False,
            False,
            None,
        ),
        ServiceMeta(
            "public_health",
            "Government Public Health Activities",
            "hce",
            2,
            False,
            True,
            120,
        ),
        ServiceMeta("investment", "Investment", "nhe", 1, True, True, 130),
        ServiceMeta("research", "Non-Commercial Research", "investment", 2, False, False, None),
        ServiceMeta(
            "structures_equipment",
            "Structures and Equipment",
            "investment",
            2,
            True,
            False,
            None,
        ),
        ServiceMeta("structures", "Structures", "structures_equipment", 3, False, False, None),
        ServiceMeta("equipment", "Equipment", "structures_equipment", 3, False, False, None),
    )
}


def funding(
    code: str,
    name: str,
    parent: str | None,
    level: int,
    *,
    subtotal: bool = False,
    additive: bool = False,
    order: int | None = None,
) -> FundingMeta:
    return FundingMeta(code, name, parent, level, subtotal, additive, order)


ALL_SOURCES = funding("all_sources", "All sources", None, 0, subtotal=True)
STANDARD_FUNDING: dict[int, FundingMeta] = {
    0: ALL_SOURCES,
    1: funding("out_of_pocket", "Out of pocket", "all_sources", 1, additive=True, order=10),
    2: funding("health_insurance", "Health Insurance", "all_sources", 1, subtotal=True),
    3: funding(
        "private_health_insurance",
        "Private Health Insurance",
        "health_insurance",
        2,
        additive=True,
        order=20,
    ),
    4: funding("medicare", "Medicare", "health_insurance", 2, additive=True, order=30),
    5: funding("medicaid", "Medicaid (Title XIX)", "health_insurance", 2, subtotal=True, additive=True, order=40),
    6: funding("medicaid_federal", "Federal", "medicaid", 3),
    7: funding("medicaid_state_local", "State and Local", "medicaid", 3),
    8: funding(
        "chip",
        "CHIP (Title XIX and Title XXI)",
        "health_insurance",
        2,
        subtotal=True,
        additive=True,
        order=50,
    ),
    9: funding("chip_federal", "Federal", "chip", 3),
    10: funding("chip_state_local", "State and Local", "chip", 3),
    11: funding(
        "department_of_defense",
        "Department of Defense",
        "health_insurance",
        2,
        additive=True,
        order=60,
    ),
    12: funding(
        "department_of_veterans_affairs",
        "Department of Veterans Affairs",
        "health_insurance",
        2,
        additive=True,
        order=70,
    ),
    13: funding(
        "other_third_party_payers_programs",
        "Other Third Party Payers and Programs",
        "all_sources",
        1,
        subtotal=True,
        additive=True,
        order=80,
    ),
    14: funding("worksite_health_care", "Worksite Health Care", "other_third_party_payers_programs", 2),
    15: funding("other_private_revenues", "Other Private Revenues", "other_third_party_payers_programs", 2),
    16: funding("indian_health_service", "Indian Health Service", "other_third_party_payers_programs", 2),
    17: funding("workers_compensation", "Workers' Compensation", "other_third_party_payers_programs", 2),
    18: funding("general_assistance", "General Assistance", "other_third_party_payers_programs", 2),
    19: funding(
        "maternal_child_health",
        "Maternal/Child Health",
        "other_third_party_payers_programs",
        2,
        subtotal=True,
    ),
    20: funding("maternal_child_health_federal", "Federal", "maternal_child_health", 3),
    21: funding("maternal_child_health_state_local", "State and Local", "maternal_child_health", 3),
    22: funding(
        "vocational_rehabilitation",
        "Vocational Rehabilitation",
        "other_third_party_payers_programs",
        2,
        subtotal=True,
    ),
    23: funding("vocational_rehabilitation_federal", "Federal", "vocational_rehabilitation", 3),
    24: funding("vocational_rehabilitation_state_local", "State and Local", "vocational_rehabilitation", 3),
    25: funding("other_federal_programs", "Other Federal Programs", "other_third_party_payers_programs", 2),
    26: funding("samhsa", "SAMHSA", "other_third_party_payers_programs", 2),
    27: funding(
        "other_state_local_programs",
        "Other State and Local Programs",
        "other_third_party_payers_programs",
        2,
    ),
    28: funding("school_health", "School Health", "other_third_party_payers_programs", 2),
    29: funding(
        "total_cms_programs",
        "Total CMS Programs (Medicaid, CHIP and Medicare)",
        None,
        0,
        subtotal=True,
    ),
}

PUBLIC_HEALTH_SOURCE = funding(
    "government_public_health_activities",
    "Government Public Health Activities",
    "all_sources",
    1,
    subtotal=True,
    additive=True,
    order=90,
)
PUBLIC_HEALTH_FEDERAL = funding("public_health_federal", "Federal", PUBLIC_HEALTH_SOURCE.code, 2)
PUBLIC_HEALTH_STATE_LOCAL = funding("public_health_state_local", "State and Local", PUBLIC_HEALTH_SOURCE.code, 2)
INVESTMENT_SOURCE = funding(
    "investment",
    "Investment",
    "all_sources",
    1,
    subtotal=True,
    additive=True,
    order=100,
)
INVESTMENT_RESEARCH = funding("investment_research", "Research", "investment", 2)
INVESTMENT_STRUCTURES_EQUIPMENT = funding(
    "investment_structures_equipment", "Structures and Equipment", "investment", 2
)
PRIVATE_FUNDS = funding("private_funds", "Private funds", "all_sources", 1)
FEDERAL_FUNDS = funding("federal_funds", "Federal funds", "all_sources", 1)
STATE_LOCAL_FUNDS = funding("state_local_funds", "State and local funds", "all_sources", 1)


BLOCKS: tuple[tuple[int, str], ...] = (
    (72, "phc"),
    (102, "hospital_care"),
    (132, "physician_clinical_services"),
    (162, "dental_services"),
    (192, "other_professional_services"),
    (222, "home_health_care"),
    (252, "other_nondurable_medical_products"),
    (282, "prescription_drugs"),
    (312, "durable_medical_equipment"),
    (342, "nursing_care_ccrc"),
    (372, "other_health_residential_personal_care"),
    (402, "administration_nonmedical_insurance"),
    (432, "state_local_administration"),
    (462, "federal_administration"),
    (492, "nonmedical_insurance"),
)


ROOT_LABELS = {
    2: "Total National Health Expenditures",
    39: "Health Consumption Expenditures",
    72: "Personal Health Care",
    102: "Total Hospital Expenditures",
    132: "Total Physician and Clinical Expenditures",
    162: "Total Dental Services Expenditures",
    192: "Total Other Professional Services Expenditures",
    222: "Total Home Health Care Expenditures",
    252: "Other Non-Durable Medical Products Expenditures",
    282: "Total Prescription Drug Expenditures",
    312: "Total Durable Medical Equipment Expenditures",
    342: "Total Nursing Care Facilities and Continuing Care Retirement Communities",
    372: "Total Other Health, Residential, and Personal Care Expenditures",
    402: "Total Administration and Total Non-Medical Insurance Expenditures",
    432: "State and Local  Administration Expenditures",
    462: "Federal Administration Expenditures",
    492: "Non-Medical Insurance Expenditures",
    522: "Public Health Activity",
    525: "Research",
    529: "Total Structures and Equipment",
    533: "     Structures",
    537: "     Equipment",
}


CANONICAL_COLUMNS = [
    "year",
    "amount_usd",
    "amount_unit",
    "service_category_code",
    "service_category_name",
    "service_category_parent",
    "funding_source_code",
    "funding_source_name",
    "sponsor_code",
    "sponsor_name",
    "program_code",
    "program_name",
    "recipient_code",
    "recipient_name",
    "claim_class",
    "accounting_view",
    "source_dataset",
    "source_file",
    "source_sheet",
    "source_row",
    "source_column",
    "source_label_raw",
    "value_status",
    "derivation",
    "source_release",
    "ingested_at",
]

STRING_COLUMNS = {
    "amount_unit",
    "service_category_code",
    "service_category_name",
    "service_category_parent",
    "funding_source_code",
    "funding_source_name",
    "sponsor_code",
    "sponsor_name",
    "program_code",
    "program_name",
    "recipient_code",
    "recipient_name",
    "claim_class",
    "accounting_view",
    "source_dataset",
    "source_file",
    "source_sheet",
    "source_column",
    "source_label_raw",
    "value_status",
    "derivation",
    "source_release",
    "ingested_at",
    "source_value_raw",
    "source_unit",
    "service_category_name_raw",
    "funding_source_parent",
    "sponsor_parent",
    "observation_role",
}
NULLABLE_INTEGER_COLUMNS = {
    "amount_usd": "Int64",
    "source_row": "Int16",
    "source_display_precision_usd": "Int64",
    "service_category_level": "Int8",
    "service_sort_order": "Int16",
    "funding_source_level": "Int8",
    "funding_source_sort_order": "Int16",
    "sponsor_level": "Int8",
}
BOOLEAN_COLUMNS = {
    "service_category_is_subtotal",
    "funding_source_is_subtotal",
    "sponsor_is_subtotal",
    "include_in_sponsor_sum",
    "include_in_nhe_service_sum",
    "include_in_service_funding_sum",
}


def enforce_dataframe_dtypes(dataframe: pd.DataFrame) -> pd.DataFrame:
    """Prevent all-null canonical fields from becoming untyped/null Parquet columns."""
    dataframe["year"] = dataframe["year"].astype("int16")
    for column in STRING_COLUMNS.intersection(dataframe.columns):
        dataframe[column] = pd.array(dataframe[column], dtype="string")
    for column, dtype in NULLABLE_INTEGER_COLUMNS.items():
        if column in dataframe:
            dataframe[column] = pd.array(dataframe[column], dtype=dtype)
    for column in BOOLEAN_COLUMNS.intersection(dataframe.columns):
        dataframe[column] = pd.array(dataframe[column], dtype="boolean")
    return dataframe


def manifest_record(filename: str) -> dict[str, Any]:
    manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    matches = [record for record in manifest["records"] if record["original_filename"] == filename]
    if len(matches) != 1:
        raise ValueError(f"Expected one manifest record for {filename}; found {len(matches)}")
    return matches[0]


def read_principal_rows(raw_dir: Path = RAW_DIR) -> list[list[str]]:
    with zipfile.ZipFile(raw_dir / PRINCIPAL_ARCHIVE) as archive:
        with archive.open(PRINCIPAL_MEMBER) as binary:
            rows = list(csv.reader(io.TextIOWrapper(binary, encoding="cp1252", newline="")))
    if len(rows) != 545 or {len(row) for row in rows} != {66}:
        raise ValueError(f"Unexpected {PRINCIPAL_MEMBER} shape: {len(rows)} rows")
    expected_years = [str(year) for year in range(1960, 2025)]
    if rows[1][1:] != expected_years:
        raise ValueError("Principal CMS year header is not exactly 1960-2024")
    for row_index, expected_label in ROOT_LABELS.items():
        if rows[row_index][0] != expected_label:
            raise ValueError(
                f"CMS row {row_index + 1} changed: expected {expected_label!r}, "
                f"found {rows[row_index][0]!r}"
            )
    template = [rows[72 + offset][0] for offset in range(1, 30)]
    for start, _ in BLOCKS[1:]:
        observed = [rows[start + offset][0] for offset in range(1, 30)]
        if observed != template:
            raise ValueError(f"Repeated funding hierarchy changed at source row {start + 1}")
    return rows


def build_row_contexts() -> dict[int, tuple[str, FundingMeta]]:
    contexts: dict[int, tuple[str, FundingMeta]] = {}

    def add(row_index: int, service_code: str, source: FundingMeta) -> None:
        if row_index in contexts:
            raise AssertionError(f"Duplicate row context: {row_index + 1}")
        contexts[row_index] = (service_code, source)

    add(2, "nhe", ALL_SOURCES)
    for offset in range(1, 29):
        add(2 + offset, "nhe", STANDARD_FUNDING[offset])
    add(31, "nhe", PUBLIC_HEALTH_SOURCE)
    add(32, "nhe", PUBLIC_HEALTH_FEDERAL)
    add(33, "nhe", PUBLIC_HEALTH_STATE_LOCAL)
    add(34, "nhe", INVESTMENT_SOURCE)
    add(35, "nhe", INVESTMENT_RESEARCH)
    add(36, "nhe", INVESTMENT_STRUCTURES_EQUIPMENT)
    add(37, "nhe", STANDARD_FUNDING[29])

    add(39, "hce", ALL_SOURCES)
    for offset in range(1, 29):
        add(39 + offset, "hce", STANDARD_FUNDING[offset])
    add(68, "hce", PUBLIC_HEALTH_SOURCE)
    add(69, "hce", PUBLIC_HEALTH_FEDERAL)
    add(70, "hce", PUBLIC_HEALTH_STATE_LOCAL)
    add(71, "hce", STANDARD_FUNDING[29])

    for start, service_code in BLOCKS:
        for offset in range(30):
            add(start + offset, service_code, STANDARD_FUNDING[offset])

    for row_index, service_code, source in (
        (522, "public_health", ALL_SOURCES),
        (523, "public_health", FEDERAL_FUNDS),
        (524, "public_health", STATE_LOCAL_FUNDS),
        (525, "research", ALL_SOURCES),
        (526, "research", PRIVATE_FUNDS),
        (527, "research", FEDERAL_FUNDS),
        (528, "research", STATE_LOCAL_FUNDS),
        (529, "structures_equipment", ALL_SOURCES),
        (530, "structures_equipment", PRIVATE_FUNDS),
        (531, "structures_equipment", FEDERAL_FUNDS),
        (532, "structures_equipment", STATE_LOCAL_FUNDS),
        (533, "structures", ALL_SOURCES),
        (534, "structures", PRIVATE_FUNDS),
        (535, "structures", FEDERAL_FUNDS),
        (536, "structures", STATE_LOCAL_FUNDS),
        (537, "equipment", ALL_SOURCES),
        (538, "equipment", PRIVATE_FUNDS),
        (539, "equipment", FEDERAL_FUNDS),
        (540, "equipment", STATE_LOCAL_FUNDS),
    ):
        add(row_index, service_code, source)

    if len(contexts) != 538:
        raise AssertionError(f"Expected 538 expenditure rows; mapped {len(contexts)}")
    return contexts


def parse_millions(raw_token: str) -> tuple[int | None, str]:
    stripped = raw_token.strip()
    if not stripped:
        return None, "structural_blank"
    if stripped in {"-", "—", "–"}:
        return None, "not_applicable"
    number = Decimal(stripped.replace(",", ""))
    if number != number.to_integral_value():
        raise ValueError(f"Expected integer millions, found {raw_token!r}")
    return int(number) * 1_000_000, "reported"


def normalize_source_service(raw_dir: Path = RAW_DIR) -> pd.DataFrame:
    rows = read_principal_rows(raw_dir)
    contexts = build_row_contexts()
    manifest = manifest_record(PRINCIPAL_ARCHIVE)
    block_root_for_row: dict[int, int] = {}
    roots = [2, 39, *[start for start, _ in BLOCKS], 522, 525, 529, 533, 537]
    sorted_roots = sorted(roots)
    for row_index in contexts:
        block_root_for_row[row_index] = max(root for root in sorted_roots if root <= row_index)

    records: list[dict[str, Any]] = []
    for row_index, (service_code, source) in sorted(contexts.items()):
        service = SERVICES[service_code]
        service_label_raw = rows[block_root_for_row[row_index]][0]
        for column_index, year in enumerate(range(1960, 2025), start=1):
            raw_token = rows[row_index][column_index]
            amount_usd, value_status = parse_millions(raw_token)
            records.append(
                {
                    "year": year,
                    "amount_usd": amount_usd,
                    "amount_unit": "USD current dollars",
                    "service_category_code": service.code,
                    "service_category_name": service.name,
                    "service_category_parent": service.parent,
                    "funding_source_code": source.code,
                    "funding_source_name": source.name,
                    "sponsor_code": None,
                    "sponsor_name": None,
                    "program_code": None,
                    "program_name": None,
                    "recipient_code": None,
                    "recipient_name": None,
                    "claim_class": None,
                    "accounting_view": "source_of_funds_by_service",
                    "source_dataset": (
                        "National Health Expenditures by type of service and source of funds, "
                        "CY 1960-2024"
                    ),
                    "source_file": f"{PRINCIPAL_ARCHIVE}::{PRINCIPAL_MEMBER}",
                    "source_sheet": PRINCIPAL_MEMBER,
                    "source_row": row_index + 1,
                    "source_column": get_column_letter(column_index + 1),
                    "source_label_raw": rows[row_index][0],
                    "value_status": value_status,
                    "derivation": None,
                    "source_release": manifest["source_release"],
                    "ingested_at": manifest["downloaded_at"],
                    "source_value_raw": raw_token,
                    "source_unit": "USD millions, current dollars",
                    "source_display_precision_usd": 1_000_000,
                    "service_category_name_raw": service_label_raw,
                    "service_category_level": service.level,
                    "service_category_is_subtotal": service.is_subtotal,
                    "service_sort_order": service.ledger_order,
                    "funding_source_parent": source.parent,
                    "funding_source_level": source.level,
                    "funding_source_is_subtotal": source.is_subtotal,
                    "funding_source_sort_order": source.ledger_order,
                    "include_in_nhe_service_sum": (
                        service.include_in_nhe_service_sum and source.code == "all_sources"
                    ),
                    "include_in_service_funding_sum": source.include_in_service_funding_sum,
                    "observation_role": (
                        "reported_total"
                        if source.code == "all_sources"
                        else "overlapping_analytical_subtotal"
                        if source.code == "total_cms_programs"
                        else "reported_subtotal"
                        if source.is_subtotal
                        else "reported_component"
                    ),
                }
            )

    dataframe = enforce_dataframe_dtypes(pd.DataFrame.from_records(records))
    return dataframe[CANONICAL_COLUMNS + [column for column in dataframe if column not in CANONICAL_COLUMNS]]


SPONSORS: dict[int, SponsorMeta] = {
    4: SponsorMeta("all_sponsors", "All sponsors", None, 0, True, False),
    5: SponsorMeta(
        "business_households_other_private",
        "Business, Households and Other Private",
        "all_sponsors",
        1,
        True,
        False,
    ),
    6: SponsorMeta("private_business", "Private Business", "business_households_other_private", 2, False, True),
    7: SponsorMeta("household", "Household", "business_households_other_private", 2, False, True),
    8: SponsorMeta(
        "other_private_sponsors",
        "Other Private Sponsors",
        "business_households_other_private",
        2,
        False,
        True,
    ),
    9: SponsorMeta("government", "Government", "all_sponsors", 1, True, False),
    10: SponsorMeta("federal_government", "Federal Government", "government", 2, False, True),
    11: SponsorMeta(
        "state_local_government", "State and Local Government", "government", 2, False, True
    ),
}


def normalize_sponsor(raw_dir: Path = RAW_DIR) -> pd.DataFrame:
    manifest = manifest_record(TABLES_ARCHIVE)
    with zipfile.ZipFile(raw_dir / TABLES_ARCHIVE) as archive:
        workbook = openpyxl.load_workbook(
            io.BytesIO(archive.read(SPONSOR_MEMBER)), read_only=True, data_only=True
        )
    if workbook.sheetnames != ["Table 5"]:
        raise ValueError(f"Unexpected sponsor workbook sheets: {workbook.sheetnames}")
    worksheet = workbook["Table 5"]
    years = [worksheet.cell(2, column).value for column in range(2, 40)]
    if years != list(range(1987, 2025)):
        raise ValueError("Sponsor table year header is not exactly 1987-2024")

    records: list[dict[str, Any]] = []
    nhe = SERVICES["nhe"]
    for source_row, sponsor in SPONSORS.items():
        source_label_raw = worksheet.cell(source_row, 1).value
        for column_index, year in enumerate(range(1987, 2025), start=2):
            raw_value = worksheet.cell(source_row, column_index).value
            if not isinstance(raw_value, (int, float)):
                raise ValueError(
                    f"Sponsor amount is not numeric at {get_column_letter(column_index)}{source_row}: "
                    f"{raw_value!r}"
                )
            amount_usd = int(Decimal(str(raw_value)) * Decimal("1000000000"))
            records.append(
                {
                    "year": year,
                    "amount_usd": amount_usd,
                    "amount_unit": "USD current dollars",
                    "service_category_code": nhe.code,
                    "service_category_name": nhe.name,
                    "service_category_parent": nhe.parent,
                    "funding_source_code": None,
                    "funding_source_name": None,
                    "sponsor_code": sponsor.code,
                    "sponsor_name": sponsor.name,
                    "program_code": None,
                    "program_name": None,
                    "recipient_code": None,
                    "recipient_name": None,
                    "claim_class": None,
                    "accounting_view": "sponsor",
                    "source_dataset": "NHE Tables, Table 5: National Health Expenditures by Type of Sponsor",
                    "source_file": f"{TABLES_ARCHIVE}::{SPONSOR_MEMBER}",
                    "source_sheet": "Table 5",
                    "source_row": source_row,
                    "source_column": get_column_letter(column_index),
                    "source_label_raw": source_label_raw,
                    "value_status": "reported",
                    "derivation": None,
                    "source_release": manifest["source_release"],
                    "ingested_at": manifest["downloaded_at"],
                    "source_value_raw": str(raw_value),
                    "source_unit": "USD billions, current dollars",
                    "source_display_precision_usd": 100_000_000,
                    "service_category_name_raw": "National Health Expenditures",
                    "service_category_level": nhe.level,
                    "service_category_is_subtotal": nhe.is_subtotal,
                    "service_sort_order": None,
                    "funding_source_parent": None,
                    "funding_source_level": None,
                    "funding_source_is_subtotal": None,
                    "funding_source_sort_order": None,
                    "sponsor_parent": sponsor.parent,
                    "sponsor_level": sponsor.level,
                    "sponsor_is_subtotal": sponsor.is_subtotal,
                    "include_in_sponsor_sum": sponsor.include_in_sponsor_sum,
                    "include_in_nhe_service_sum": False,
                    "include_in_service_funding_sum": False,
                    "observation_role": "reported_subtotal" if sponsor.is_subtotal else "reported_component",
                }
            )
    dataframe = enforce_dataframe_dtypes(pd.DataFrame.from_records(records))
    return dataframe[CANONICAL_COLUMNS + [column for column in dataframe if column not in CANONICAL_COLUMNS]]


def write_duckdb(source_service_path: Path, sponsor_path: Path, database_path: Path) -> None:
    database_path.parent.mkdir(parents=True, exist_ok=True)
    with duckdb.connect(str(database_path)) as connection:
        connection.execute(
            "CREATE OR REPLACE TABLE nhea_source_service AS SELECT * FROM read_parquet(?)",
            [str(source_service_path)],
        )
        connection.execute(
            "CREATE OR REPLACE TABLE nhea_sponsor AS SELECT * FROM read_parquet(?)",
            [str(sponsor_path)],
        )
        connection.execute(
            """
            CREATE OR REPLACE VIEW nhea_source_service_reported AS
            SELECT * FROM nhea_source_service WHERE value_status = 'reported'
            """
        )
        connection.execute(
            """
            CREATE OR REPLACE VIEW sponsor_leaf_totals AS
            SELECT * FROM nhea_sponsor WHERE include_in_sponsor_sum
            """
        )


def normalize_all() -> tuple[pd.DataFrame, pd.DataFrame]:
    SOURCE_SERVICE_PATH.parent.mkdir(parents=True, exist_ok=True)
    source_service = normalize_source_service()
    sponsor = normalize_sponsor()
    source_service.to_parquet(SOURCE_SERVICE_PATH, index=False)
    sponsor.to_parquet(SPONSOR_PATH, index=False)
    write_duckdb(SOURCE_SERVICE_PATH, SPONSOR_PATH, DUCKDB_PATH)
    return source_service, sponsor


def main() -> None:
    source_service, sponsor = normalize_all()
    print(f"Wrote {len(source_service):,} records to {SOURCE_SERVICE_PATH}")
    print(f"Wrote {len(sponsor):,} records to {SPONSOR_PATH}")
    print(f"Loaded both accounting views into {DUCKDB_PATH}")


if __name__ == "__main__":
    main()
