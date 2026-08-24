"""Normalize official CMS payment files without retaining copyrighted descriptions."""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import load_workbook

from reconciliation_healthcare.paths import RULEBOOK_MANIFEST_PATH, RULEBOOK_RAW_DIR


PARAMETER_COLUMNS = [
    "parameter_id",
    "rule_id",
    "parameter_name",
    "effective_start",
    "effective_end",
    "key_type",
    "key_value",
    "numeric_value",
    "string_value",
    "unit",
    "geography",
    "provider_id",
    "code",
    "setting",
    "source_artifact_id",
    "source_locator",
    "value_status",
]

ASSIGNMENT_COLUMNS = [
    "assignment_id",
    "rule_id",
    "payment_system",
    "assignment_type",
    "code_system",
    "code",
    "modifier",
    "effective_start",
    "effective_end",
    "status_code",
    "status_indicator",
    "apc",
    "work_rvu",
    "nonfacility_pe_rvu",
    "nonfacility_na_indicator",
    "facility_pe_rvu",
    "facility_na_indicator",
    "malpractice_rvu",
    "multiple_procedure_indicator",
    "bilateral_surgery_indicator",
    "assistant_surgery_indicator",
    "co_surgery_indicator",
    "team_surgery_indicator",
    "relative_weight",
    "payment_rate",
    "source_artifact_id",
    "source_locator",
    "value_status",
]

PFS_ASSIGNMENT_VALUE_COLUMNS = (
    "status_code",
    "work_rvu",
    "nonfacility_pe_rvu",
    "nonfacility_na_indicator",
    "facility_pe_rvu",
    "facility_na_indicator",
    "malpractice_rvu",
    "multiple_procedure_indicator",
    "bilateral_surgery_indicator",
    "assistant_surgery_indicator",
    "co_surgery_indicator",
    "team_surgery_indicator",
)


def _date(value: str) -> date:
    return date.fromisoformat(value)


def _clean_string(value: Any) -> str | None:
    if value is None:
        return None
    cleaned = str(value).strip()
    return cleaned or None


def _decimal(value: Any) -> Decimal | None:
    cleaned = _clean_string(value)
    if cleaned is None or cleaned in {".", "-", "—", "N/A", "NA"}:
        return None
    cleaned = cleaned.replace("$", "").replace(",", "").replace("%", "")
    try:
        return Decimal(cleaned)
    except InvalidOperation as error:
        raise ValueError(f"Expected numeric CMS value, received {value!r}") from error


def _decode_csv(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("cms-csv", raw, 0, 1, "unsupported source encoding")


def _member(zip_file: zipfile.ZipFile, pattern: str) -> str:
    matches = [name for name in zip_file.namelist() if re.search(pattern, name, re.I)]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one ZIP member matching {pattern!r}; found {len(matches)}: {matches}"
        )
    return matches[0]


def _parameter(**values: Any) -> dict[str, Any]:
    return {column: values.get(column) for column in PARAMETER_COLUMNS}


def _assignment(**values: Any) -> dict[str, Any]:
    return {column: values.get(column) for column in ASSIGNMENT_COLUMNS}


def _artifact_records(manifest: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {record["source_artifact_id"]: record for record in manifest["records"]}


def _pfs_release(artifact_id: str) -> str:
    return artifact_id.removeprefix("cms_pfs_")


def _parse_pfs_snapshot(
    raw_dir: Path,
    record: dict[str, Any],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    artifact_id = record["source_artifact_id"]
    version = _pfs_release(artifact_id)
    path = raw_dir / record["relative_path"]
    assignments: list[dict[str, Any]] = []
    parameters: list[dict[str, Any]] = []
    with zipfile.ZipFile(path) as archive:
        rvu_member = _member(archive, r"PPRRVU.*\.csv$")
        rows = list(csv.reader(io.StringIO(_decode_csv(archive.read(rvu_member)))))
        header_index = next(
            index for index, row in enumerate(rows) if row and row[0].strip() == "HCPCS"
        )
        conversion_factors: set[Decimal] = set()
        for source_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            if not row or not _clean_string(row[0]) or len(row) < 31:
                continue
            code = row[0].strip().upper()
            if not re.fullmatch(r"[A-Z0-9]{5}", code):
                continue
            modifier = (_clean_string(row[1]) or "").upper()
            status = (_clean_string(row[3]) or "").upper()
            conversion_factor = _decimal(row[24])
            if status == "A" and conversion_factor is not None:
                conversion_factors.add(conversion_factor)
            assignments.append(
                _assignment(
                    assignment_id=f"pfs.{version}.{code}.{modifier or 'blank'}",
                    rule_id=f"pfs.{version}.base_payment",
                    payment_system="PFS",
                    assignment_type="pfs_rvu",
                    code_system="HCPCS",
                    code=code,
                    modifier=modifier,
                    effective_start=_date(record["effective_start"]),
                    effective_end=_date(record["effective_end"]),
                    status_code=status,
                    work_rvu=_decimal(row[5]),
                    nonfacility_pe_rvu=_decimal(row[6]),
                    nonfacility_na_indicator=_clean_string(row[7]),
                    facility_pe_rvu=_decimal(row[8]),
                    facility_na_indicator=_clean_string(row[9]),
                    malpractice_rvu=_decimal(row[10]),
                    multiple_procedure_indicator=_clean_string(row[18]),
                    bilateral_surgery_indicator=_clean_string(row[19]),
                    assistant_surgery_indicator=_clean_string(row[20]),
                    co_surgery_indicator=_clean_string(row[21]),
                    team_surgery_indicator=_clean_string(row[22]),
                    source_artifact_id=artifact_id,
                    source_locator=f"{rvu_member} row {source_index}",
                    value_status="published",
                )
            )
        if len(conversion_factors) != 1:
            raise ValueError(
                f"Expected one active-code conversion factor in {artifact_id}; "
                f"found {sorted(conversion_factors)}"
            )
        conversion_factor = next(iter(conversion_factors))
        parameters.append(
            _parameter(
                parameter_id=f"pfs.{version}.conversion_factor",
                rule_id=f"pfs.{version}.conversion_factor",
                parameter_name="conversion_factor",
                effective_start=_date(record["effective_start"]),
                effective_end=_date(record["effective_end"]),
                key_type="national",
                key_value="US",
                numeric_value=conversion_factor,
                unit="USD per geographically adjusted RVU",
                source_artifact_id=artifact_id,
                source_locator=f"{rvu_member} Conversion Factor column; active-code value",
                value_status="published",
            )
        )

        gpci_member = _member(archive, r"(^|/)GPCI2024\.csv$")
        gpci_rows = list(csv.reader(io.StringIO(_decode_csv(archive.read(gpci_member)))))
        gpci_header = next(
            index
            for index, row in enumerate(gpci_rows)
            if row and row[0].strip().startswith("Medicare Administrative Contractor")
        )
        seen_localities: set[str] = set()
        for source_index, row in enumerate(
            gpci_rows[gpci_header + 1 :], start=gpci_header + 2
        ):
            if len(row) < 7:
                continue
            mac = _clean_string(row[0])
            state = _clean_string(row[1])
            locality_number = _clean_string(row[2])
            if not (mac and state and locality_number and len(state) == 2):
                continue
            locality_number = locality_number.zfill(2)
            locality = f"{state.upper()}:{locality_number}"
            if locality in seen_localities:
                raise ValueError(f"Duplicate canonical PFS locality {locality} in {artifact_id}")
            seen_localities.add(locality)
            for column_index, parameter_name, source_label in (
                (4, "work_gpci", "2024 PW GPCI"),
                (5, "pe_gpci", "2024 PE GPCI"),
                (6, "mp_gpci", "2024 MP GPCI"),
            ):
                parameters.append(
                    _parameter(
                        parameter_id=f"pfs.{version}.{locality}.{parameter_name}",
                        rule_id=f"pfs.{version}.gpci",
                        parameter_name=parameter_name,
                        effective_start=_date(record["effective_start"]),
                        effective_end=_date(record["effective_end"]),
                        key_type="state_locality",
                        key_value=locality,
                        numeric_value=_decimal(row[column_index]),
                        unit="index",
                        geography=locality,
                        string_value=mac.zfill(5),
                        source_artifact_id=artifact_id,
                        source_locator=(
                            f"{gpci_member} row {source_index}, column {source_label}"
                        ),
                        value_status="published",
                    )
                )
    return assignments, parameters


def _mark_pfs_effective_date_uncertainty(
    snapshots: list[list[dict[str, Any]]],
) -> None:
    """Do not pretend every changed quarterly row took effect on quarter day one."""
    known_quarter_boundaries = {
        ("rvu24b", "J0576"),
        ("rvu24c", "G9037"),
        ("rvu24c", "G9038"),
        ("rvu24d", "J1170"),
    }
    prior: dict[tuple[str, str], tuple[Any, ...]] = {}
    for snapshot_index, rows in enumerate(snapshots):
        current: dict[tuple[str, str], tuple[Any, ...]] = {}
        for row in rows:
            key = (row["code"], row["modifier"])
            values = tuple(row[column] for column in PFS_ASSIGNMENT_VALUE_COLUMNS)
            current[key] = values
            if snapshot_index:
                changed = prior.get(key) != values
                version = row["assignment_id"].split(".")[1]
                if changed and (version, row["code"]) not in known_quarter_boundaries:
                    row["value_status"] = "unresolved_effective_date"
                    row["source_locator"] += (
                        "; quarterly snapshot row changed and no row-specific CR date was normalized"
                    )
                elif changed:
                    row["value_status"] = "published_cr_effective_date"
        prior = current


def normalize_pfs(
    raw_dir: Path,
    artifacts: dict[str, dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    snapshots: list[list[dict[str, Any]]] = []
    parameters: list[dict[str, Any]] = []
    for artifact_id in (
        "cms_pfs_rvu24a",
        "cms_pfs_rvu24ar",
        "cms_pfs_rvu24b",
        "cms_pfs_rvu24c",
        "cms_pfs_rvu24d",
    ):
        assignments, release_parameters = _parse_pfs_snapshot(
            raw_dir, artifacts[artifact_id]
        )
        snapshots.append(assignments)
        parameters.extend(release_parameters)
    _mark_pfs_effective_date_uncertainty(snapshots)
    return [row for snapshot in snapshots for row in snapshot], parameters


def _opps_rows(raw_dir: Path, record: dict[str, Any]) -> list[dict[str, Any]]:
    artifact_id = record["source_artifact_id"]
    quarter = artifact_id.split("_")[3]
    path = raw_dir / record["relative_path"]
    rows_out: list[dict[str, Any]] = []
    with zipfile.ZipFile(path) as archive:
        csv_member = _member(archive, r"\.csv$")
        rows = list(csv.reader(io.StringIO(_decode_csv(archive.read(csv_member)))))
        header_index = next(
            index
            for index, row in enumerate(rows)
            if row and row[0].strip() == "HCPCS Code"
        )
        header = {value.strip().lower(): index for index, value in enumerate(rows[header_index])}
        status_index = next(
            index
            for key, index in header.items()
            if key.replace(" ", "") in {"si", "statusindicator"}
        )
        apc_index = next(index for key, index in header.items() if key.strip() == "apc")
        relative_weight_index = next(
            index for key, index in header.items() if key == "relative weight"
        )
        payment_rate_index = next(
            index for key, index in header.items() if key == "payment rate"
        )
        for source_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            if not row or not _clean_string(row[0]):
                continue
            code = row[0].strip().upper()
            if not re.fullmatch(r"[A-Z0-9]{5}", code):
                continue
            rate = _decimal(row[payment_rate_index])
            status = (_clean_string(row[status_index]) or "").upper()
            value_status = "published_quarter_snapshot"
            if status in {"G", "K"} or (rate is not None and -rate.as_tuple().exponent > 2):
                value_status = "requires_restated_drug_overlay"
            rows_out.append(
                _assignment(
                    assignment_id=f"opps.2024_{quarter}.{code}",
                    rule_id=f"opps.2024_{quarter}.hcpcs_status_apc",
                    payment_system="OPPS",
                    assignment_type="opps_hcpcs",
                    code_system="HCPCS",
                    code=code,
                    modifier="",
                    effective_start=_date(record["effective_start"]),
                    effective_end=_date(record["effective_end"]),
                    status_indicator=status,
                    apc=_clean_string(row[apc_index]),
                    relative_weight=_decimal(row[relative_weight_index]),
                    payment_rate=rate,
                    source_artifact_id=artifact_id,
                    source_locator=f"{csv_member} row {source_index}",
                    value_status=value_status,
                )
            )
    return rows_out


def _mark_opps_effective_date_uncertainty(
    snapshots: list[list[dict[str, Any]]],
) -> None:
    known_clean_boundaries = {
        ("q3", "C1761"),
        ("q4", "C1831"),
        ("q4", "G0012"),
    }
    prior: dict[str, tuple[Any, ...]] = {}
    for index, rows in enumerate(snapshots):
        current: dict[str, tuple[Any, ...]] = {}
        for row in rows:
            values = (
                row["status_indicator"],
                row["apc"],
                row["relative_weight"],
                row["payment_rate"],
            )
            current[row["code"]] = values
            if index and prior.get(row["code"]) != values:
                quarter = row["assignment_id"].split(".")[1].split("_")[1]
                if (quarter, row["code"]) in known_clean_boundaries:
                    row["value_status"] = "published_cr_effective_date"
                else:
                    row["value_status"] = "unresolved_effective_date"
                    row["source_locator"] += (
                        "; changed quarterly row without a normalized code-specific update date"
                    )
        prior = current

    # MM13568 corrected this Q1 assignment retroactive to 2024-01-01. CMS's
    # April Addendum B carries the corrected APC/rate, so copy only those public
    # numeric fields and point provenance to both official artifacts.
    q1 = next(rows for rows in snapshots if rows[0]["assignment_id"].startswith("opps.2024_q1"))
    q2 = next(rows for rows in snapshots if rows[0]["assignment_id"].startswith("opps.2024_q2"))
    q1_by_code = {row["code"]: row for row in q1}
    q2_by_code = {row["code"]: row for row in q2}
    if "C9790" in q1_by_code and "C9790" in q2_by_code:
        target = q1_by_code["C9790"]
        corrected = q2_by_code["C9790"]
        for column in ("status_indicator", "apc", "relative_weight", "payment_rate"):
            target[column] = corrected[column]
        target["value_status"] = "published_retroactive_correction"
        target["source_locator"] = (
            f"{target['source_locator']}; MM13568 retroactive correction effective 2024-01-01; "
            f"corrected numeric fields confirmed at {corrected['source_locator']}"
        )
        corrected["value_status"] = "published_quarter_snapshot"


def normalize_opps(
    raw_dir: Path,
    artifacts: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    snapshots = [
        _opps_rows(raw_dir, artifacts[artifact_id])
        for artifact_id in (
            "cms_opps_2024_q1_addendum_b",
            "cms_opps_2024_q2_addendum_b",
            "cms_opps_2024_q3_addendum_b",
            "cms_opps_2024_q4_addendum_b",
        )
    ]
    _mark_opps_effective_date_uncertainty(snapshots)
    return [row for snapshot in snapshots for row in snapshot]


IPPS_CATEGORIES = (
    "quality_submitted_ehr_user",
    "quality_submitted_not_ehr_user",
    "quality_not_submitted_ehr_user",
    "quality_not_submitted_not_ehr_user",
)


def _xlsx_from_zip(path: Path, member_pattern: str) -> tuple[Any, str]:
    with zipfile.ZipFile(path) as archive:
        member = _member(archive, member_pattern)
        workbook = load_workbook(io.BytesIO(archive.read(member)), read_only=True, data_only=True)
    return workbook, member


def _nested_xlsx_from_zip(
    path: Path,
    nested_pattern: str,
    member_pattern: str,
) -> tuple[Any, str]:
    with zipfile.ZipFile(path) as outer:
        nested_name = _member(outer, nested_pattern)
        with zipfile.ZipFile(io.BytesIO(outer.read(nested_name))) as nested:
            workbook_name = _member(nested, member_pattern)
            workbook = load_workbook(
                io.BytesIO(nested.read(workbook_name)), read_only=True, data_only=True
            )
    return workbook, f"{nested_name} > {workbook_name}"


def _ipps_standardized_amounts(
    raw_dir: Path,
    record: dict[str, Any],
    *,
    fiscal_year: int,
) -> list[dict[str, Any]]:
    path = raw_dir / record["relative_path"]
    workbook, member = _xlsx_from_zip(path, r"\.xlsx$")
    sheet_name = (
        "FY 2024 FINAL Table 1A-1E"
        if fiscal_year == 2024
        else "FY 2025 IFC Table 1A-1E"
    )
    sheet = workbook[sheet_name]
    version = f"fy{fiscal_year}"
    parameters: list[dict[str, Any]] = []
    for row_number, labor_share in ((5, "67.6"), (10, "62.0")):
        for category_index, category in enumerate(IPPS_CATEGORIES):
            for offset, parameter_name, label in (
                (0, "standardized_labor_amount", "labor"),
                (1, "standardized_nonlabor_amount", "nonlabor"),
            ):
                column = category_index * 2 + offset + 1
                cell = sheet.cell(row=row_number, column=column)
                parameters.append(
                    _parameter(
                        parameter_id=(
                            f"ipps.{version}.{category}.labor_share_{labor_share}.{label}"
                        ),
                        rule_id=f"ipps.{version}.standardized_amount",
                        parameter_name=parameter_name,
                        effective_start=_date(record["effective_start"]),
                        effective_end=_date(record["effective_end"]),
                        key_type="hospital_update_category_and_labor_share",
                        key_value=f"{category}|labor_share={labor_share}",
                        numeric_value=_decimal(cell.value),
                        unit="USD",
                        setting="acute_ipps_full_federal_rate",
                        source_artifact_id=record["source_artifact_id"],
                        source_locator=f"{member} > {sheet_name}!{cell.coordinate}",
                        value_status="published_ifc"
                        if fiscal_year == 2025
                        else "published_final",
                    )
                )
    return parameters


def _ipps_wage_parameters(
    raw_dir: Path,
    record: dict[str, Any],
    *,
    fiscal_year: int,
) -> list[dict[str, Any]]:
    path = raw_dir / record["relative_path"]
    if fiscal_year == 2024:
        workbook, member = _xlsx_from_zip(path, r"\.xlsx$")
        sheet_name = "Table 2 CN"
    else:
        workbook, member = _nested_xlsx_from_zip(
            path, r"IFC.*\.zip$", r"IFC.*\.xlsx$"
        )
        sheet_name = "Table 2 IFC"
    sheet = workbook[sheet_name]
    parameters: list[dict[str, Any]] = []
    version = f"fy{fiscal_year}"
    for source_row, values in enumerate(
        sheet.iter_rows(min_row=3, values_only=True), start=3
    ):
        ccn = _clean_string(values[0] if values else None)
        if not ccn or not re.fullmatch(r"\d{1,6}", ccn):
            continue
        ccn = ccn.zfill(6)
        if fiscal_year == 2024:
            wage = _decimal(values[6])
            transition = None
            source_column = "G"
        else:
            wage = _decimal(values[5])
            transition = _decimal(values[6])
            source_column = "F"
        if wage is None:
            continue
        payment_cbsa = _clean_string(values[12])
        parameters.append(
            _parameter(
                parameter_id=f"ipps.{version}.{ccn}.wage_index",
                rule_id=f"ipps.{version}.wage_index",
                parameter_name="wage_index",
                effective_start=_date(record["effective_start"]),
                effective_end=_date(record["effective_end"]),
                key_type="ccn",
                key_value=ccn,
                numeric_value=wage,
                string_value="with_cap",
                unit="index",
                geography=payment_cbsa,
                provider_id=ccn,
                setting="acute_ipps",
                source_artifact_id=record["source_artifact_id"],
                source_locator=f"{member} > {sheet_name}!{source_column}{source_row}",
                value_status="published_ifc"
                if fiscal_year == 2025
                else "published_correction_notice",
            )
        )
        if transition is not None:
            parameters.append(
                _parameter(
                    parameter_id=f"ipps.{version}.{ccn}.transitional_exception_wage_factor",
                    rule_id=f"ipps.{version}.wage_index",
                    parameter_name="transitional_exception_wage_factor",
                    effective_start=_date(record["effective_start"]),
                    effective_end=_date(record["effective_end"]),
                    key_type="ccn",
                    key_value=ccn,
                    numeric_value=transition,
                    unit="index-equivalent payment factor",
                    geography=payment_cbsa,
                    provider_id=ccn,
                    setting="acute_ipps",
                    source_artifact_id=record["source_artifact_id"],
                    source_locator=f"{member} > {sheet_name}!G{source_row}",
                    value_status="lookup_only_separate_payment_exception",
                )
            )
        out_migration = _decimal(values[16])
        if out_migration not in (None, Decimal("0")):
            parameters.append(
                _parameter(
                    parameter_id=f"ipps.{version}.{ccn}.out_migration_adjustment",
                    rule_id=f"ipps.{version}.wage_index",
                    parameter_name="out_migration_adjustment",
                    effective_start=_date(record["effective_start"]),
                    effective_end=_date(record["effective_end"]),
                    key_type="ccn",
                    key_value=ccn,
                    numeric_value=out_migration,
                    unit="index increment",
                    geography=payment_cbsa,
                    provider_id=ccn,
                    setting="acute_ipps",
                    source_artifact_id=record["source_artifact_id"],
                    source_locator=f"{member} > {sheet_name}!Q{source_row}",
                    value_status="documented_not_executed",
                )
            )
    return parameters


def _ipps_drg_assignments(
    raw_dir: Path,
    record: dict[str, Any],
    *,
    fiscal_year: int,
) -> list[dict[str, Any]]:
    path = raw_dir / record["relative_path"]
    workbook, member = _xlsx_from_zip(path, r"\.xlsx$")
    sheet_name = "FY 2024 Table 5 FR" if fiscal_year == 2024 else "FY 2025 Table 5 CN"
    sheet = workbook[sheet_name]
    version = f"fy{fiscal_year}"
    assignments: list[dict[str, Any]] = []
    for source_row, values in enumerate(
        sheet.iter_rows(min_row=3, values_only=True), start=3
    ):
        raw_code = _clean_string(values[0] if values else None)
        if not raw_code or not re.fullmatch(r"\d{1,3}", raw_code):
            continue
        code = raw_code.zfill(3)
        weight = _decimal(values[7])
        if weight is None:
            continue
        assignments.append(
            _assignment(
                assignment_id=f"ipps.{version}.ms_drg.{code}",
                rule_id=f"ipps.{version}.ms_drg_weight",
                payment_system="IPPS",
                assignment_type="ms_drg_weight",
                code_system="MS-DRG",
                code=code,
                modifier="",
                effective_start=_date(record["effective_start"]),
                effective_end=_date(record["effective_end"]),
                relative_weight=weight,
                source_artifact_id=record["source_artifact_id"],
                source_locator=f"{member} > {sheet_name}!H{source_row}",
                value_status="published_correction_notice"
                if fiscal_year == 2025
                else "published_final",
            )
        )
    return assignments


def normalize_ipps(
    raw_dir: Path,
    artifacts: dict[str, dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    assignments: list[dict[str, Any]] = []
    parameters: list[dict[str, Any]] = []
    for fiscal_year in (2024, 2025):
        prefix = f"cms_ipps_fy{fiscal_year}"
        parameters.extend(
            _ipps_standardized_amounts(
                raw_dir, artifacts[f"{prefix}_table1"], fiscal_year=fiscal_year
            )
        )
        parameters.extend(
            _ipps_wage_parameters(
                raw_dir, artifacts[f"{prefix}_wage_tables"], fiscal_year=fiscal_year
            )
        )
        assignments.extend(
            _ipps_drg_assignments(
                raw_dir, artifacts[f"{prefix}_table5"], fiscal_year=fiscal_year
            )
        )
    return assignments, parameters


def _frame(records: Iterable[dict[str, Any]], columns: list[str]) -> pd.DataFrame:
    frame = pd.DataFrame.from_records(records, columns=columns)
    for column in frame.columns:
        if column in {"numeric_value", "work_rvu", "nonfacility_pe_rvu", "facility_pe_rvu", "malpractice_rvu", "relative_weight", "payment_rate"}:
            continue
        if column in {"effective_start", "effective_end"}:
            continue
        frame[column] = pd.array(frame[column], dtype="string")
    return frame


def normalize_sources(
    raw_dir: Path = RULEBOOK_RAW_DIR,
    manifest_path: Path = RULEBOOK_MANIFEST_PATH,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    artifacts = _artifact_records(manifest)
    pfs_assignments, pfs_parameters = normalize_pfs(raw_dir, artifacts)
    opps_assignments = normalize_opps(raw_dir, artifacts)
    ipps_assignments, ipps_parameters = normalize_ipps(raw_dir, artifacts)

    parameters = _frame(pfs_parameters + ipps_parameters, PARAMETER_COLUMNS)
    assignments = _frame(
        pfs_assignments + opps_assignments + ipps_assignments, ASSIGNMENT_COLUMNS
    )
    source_artifacts = pd.DataFrame.from_records(manifest["records"])
    for column in ("effective_start", "effective_end"):
        source_artifacts[column] = pd.to_datetime(source_artifacts[column]).dt.date
    source_artifacts["bytes"] = pd.array(source_artifacts["bytes"], dtype="Int64")
    for column in source_artifacts.columns:
        if column not in {"effective_start", "effective_end", "bytes"}:
            source_artifacts[column] = pd.array(source_artifacts[column], dtype="string")

    parameters = parameters.sort_values("parameter_id", ignore_index=True)
    assignments = assignments.sort_values("assignment_id", ignore_index=True)
    source_artifacts = source_artifacts.sort_values("source_artifact_id", ignore_index=True)
    return parameters, assignments, source_artifacts
