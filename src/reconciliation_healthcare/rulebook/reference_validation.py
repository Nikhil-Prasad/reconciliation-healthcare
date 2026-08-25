"""Expanded independent reference validation against pinned CMS source outputs.

The builders in this module read the authoritative raw files directly.  They do
not use normalized rulebook values to construct expected amounts, so comparisons
with the execution engines exercise both normalization and calculation.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable, Mapping, Sequence

from openpyxl import load_workbook

from reconciliation_healthcare.paths import RULEBOOK_MANIFEST_PATH, RULEBOOK_RAW_DIR
from reconciliation_healthcare.rulebook.ipps import calculate_ipps_base_payment
from reconciliation_healthcare.rulebook.models import CalculationStatus
from reconciliation_healthcare.rulebook.pfs import calculate_pfs
from reconciliation_healthcare.rulebook.store import RulebookStore
from reconciliation_healthcare.rulebook.trace import validate_trace


_CENTS = Decimal("0.01")
_NINE_PLACES = Decimal("0.000000001")


@dataclass(frozen=True)
class PFSCarrierReferenceCase:
    case_id: str
    code: str
    service_date: date
    rule_version: str
    locality: str
    setting: str
    expected_amount: Decimal
    reference_artifact_id: str
    source_member: str
    source_locator: str


@dataclass(frozen=True)
class IPPSReferenceCase:
    case_id: str
    fiscal_year: str
    ms_drg: str
    provider_ccn: str
    geography: str
    discharge_date: date
    quality_submitted: bool
    meaningful_ehr_user: bool
    relative_weight: Decimal
    wage_index: Decimal
    labor_share: Decimal
    labor_amount: Decimal
    nonlabor_amount: Decimal
    unrounded_amount: Decimal
    amount_nine_places: Decimal
    expected_amount: Decimal
    source_artifact_ids: tuple[str, ...]
    source_locators: tuple[str, ...]


@dataclass(frozen=True)
class ReferenceValidationSummary:
    payment_system: str
    matched_cases: int
    codes: tuple[str, ...]
    geographies: tuple[str, ...]
    periods: tuple[str, ...]
    settings: tuple[str, ...] = ()
    labor_share_branches: tuple[str, ...] = ()
    source_artifact_ids: tuple[str, ...] = ()


@dataclass(frozen=True)
class _CarrierLocality:
    canonical_locality: str
    member: str
    carrier: str
    carrier_locality: str


_PFS_PERIODS = (
    (
        "rvu24a",
        date(2024, 1, 15),
        "cms_pfs_carrier_2024_jan_mar8",
    ),
    (
        "rvu24ar",
        date(2024, 3, 9),
        "cms_pfs_carrier_2024_mar9_dec31",
    ),
    (
        "rvu24b",
        date(2024, 4, 1),
        "cms_pfs_carrier_2024_mar9_dec31",
    ),
    (
        "rvu24c",
        date(2024, 7, 1),
        "cms_pfs_carrier_2024_mar9_dec31",
    ),
    (
        "rvu24d",
        date(2024, 10, 1),
        "cms_pfs_carrier_2024_mar9_dec31",
    ),
)

_PFS_LOCALITIES = (
    _CarrierLocality("AL:00", "PFAL24A.TXT", "10112", "00"),
    _CarrierLocality("NY:01", "PFNY224R.TXT", "13202", "01"),
    _CarrierLocality("CA:05", "PFCA24A.TXT", "01112", "05"),
    _CarrierLocality("TX:09", "PFTX24A.TXT", "04412", "09"),
    _CarrierLocality("AK:01", "PFAK24A.TXT", "02102", "01"),
)

# The selector takes one stable-hash winner from each band after intersecting
# eligibility across all five raw PPRRVU snapshots and carrier-output coverage.
# The bands broaden the fixture without depending on licensed descriptions.
_PFS_CODE_BANDS = (
    (80000, 84999),
    (85000, 89999),
    (90000, 94999),
    (95000, 99999),
)
_PFS_SELECTION_SALT = "stage2a1-pfs-reference-v1"
_PFS_RVU_ARTIFACT_IDS = tuple(
    f"cms_pfs_{version}" for version, _, _ in _PFS_PERIODS
)


def _decimal(value: object, *, label: str) -> Decimal:
    if value is None or not str(value).strip():
        raise ValueError(f"Missing numeric {label}")
    normalized = str(value).strip().replace("$", "").replace(",", "")
    try:
        result = Decimal(normalized)
    except InvalidOperation as error:
        raise ValueError(f"Invalid numeric {label}: {value!r}") from error
    if not result.is_finite():
        raise ValueError(f"Non-finite numeric {label}: {value!r}")
    return result


def _optional_decimal(value: object) -> Decimal | None:
    if value is None or not str(value).strip():
        return None
    return _decimal(value, label="optional source value")


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _pinned_paths(
    artifact_ids: Iterable[str],
    *,
    raw_dir: Path,
    manifest_path: Path,
) -> dict[str, Path]:
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    records = {
        str(record["source_artifact_id"]): record for record in manifest["records"]
    }
    base = raw_dir.resolve()
    result: dict[str, Path] = {}
    for artifact_id in sorted(set(artifact_ids)):
        if artifact_id not in records:
            raise ValueError(f"Pinned source is absent from manifest: {artifact_id}")
        record = records[artifact_id]
        path = (raw_dir / str(record["relative_path"])).resolve()
        if not path.is_relative_to(base):
            raise ValueError(f"Pinned source escapes raw directory: {artifact_id}")
        if not path.is_file():
            raise FileNotFoundError(path)
        if path.stat().st_size != int(record["bytes"]):
            raise ValueError(f"Pinned byte count mismatch: {artifact_id}")
        if _sha256(path) != str(record["sha256"]):
            raise ValueError(f"Pinned SHA-256 mismatch: {artifact_id}")
        result[artifact_id] = path
    return result


def _positive_decimal(value: object) -> bool:
    parsed = _optional_decimal(value)
    return parsed is not None and parsed > 0


def _zero_decimal(value: object) -> bool:
    parsed = _optional_decimal(value)
    return parsed is not None and parsed == 0


def _eligible_pfs_codes(path: Path) -> set[str]:
    """Read uncomplicated blank-modifier codes from one raw PPRRVU snapshot."""

    with zipfile.ZipFile(path) as archive:
        members = [
            name
            for name in archive.namelist()
            if "PPRRVU" in name.upper() and name.lower().endswith(".csv")
        ]
        if len(members) != 1:
            raise ValueError(
                f"Expected one PPRRVU CSV in {path.name}; found {members}"
            )
        with archive.open(members[0]) as raw:
            rows = list(
                csv.reader(
                    io.TextIOWrapper(
                        raw,
                        encoding="utf-8-sig",
                        errors="replace",
                        newline="",
                    )
                )
            )

    header_index = next(
        (index for index, row in enumerate(rows) if row and _text(row[0]) == "HCPCS"),
        None,
    )
    if header_index is None:
        raise ValueError(f"PPRRVU header is absent from {path.name}")

    eligible: set[str] = set()
    seen_blank_modifier: set[str] = set()
    for row in rows[header_index + 1 :]:
        if len(row) < 31:
            continue
        code = _text(row[0]).upper()
        modifier = _text(row[1]).upper()
        if len(code) != 5 or not code.isdigit() or modifier:
            continue
        if code in seen_blank_modifier:
            raise ValueError(f"Duplicate blank-modifier PPRRVU row in {path.name}: {code}")
        seen_blank_modifier.add(code)

        # These criteria deliberately exclude every code-level branch that the
        # Stage 2A base calculator labels unsupported or does not execute.
        if _text(row[3]).upper() != "A":
            continue
        if not all(_positive_decimal(row[index]) for index in (5, 6, 8, 10)):
            continue
        if _text(row[7]) or _text(row[9]):
            continue
        if _text(row[13]) != "0" or _text(row[14]).upper() != "XXX":
            continue
        if not all(_zero_decimal(row[index]) for index in (15, 16, 17)):
            continue
        if not all(_text(row[index]) in {"0", "9"} for index in range(18, 23)):
            continue
        if _text(row[23]):
            continue
        if (
            _text(row[25]) != "09"
            or _text(row[26]) != "0"
            or _text(row[27]) != "99"
        ):
            continue
        if not all(_zero_decimal(row[index]) for index in (28, 29, 30)):
            continue
        eligible.add(code)

    if not eligible:
        raise ValueError(f"No uncomplicated PFS codes found in {path.name}")
    return eligible


def _carrier_reference_codes(
    path: Path,
    locality: _CarrierLocality,
    candidates: set[str],
) -> set[str]:
    """Return candidate codes with positive facility/NF output rows."""

    matches: set[str] = set()
    with zipfile.ZipFile(path) as archive:
        if locality.member not in archive.namelist():
            raise ValueError(
                f"Carrier member {locality.member!r} is absent from {path.name}"
            )
        with archive.open(locality.member) as raw:
            rows = csv.reader(io.TextIOWrapper(raw, encoding="ascii", newline=""))
            for row in rows:
                if (
                    len(row) < 7
                    or _text(row[0]) != "2024"
                    or _text(row[1]) != locality.carrier
                    or _text(row[2]) != locality.carrier_locality
                    or _text(row[4])
                ):
                    continue
                code = _text(row[3]).upper()
                if code not in candidates:
                    continue
                if code in matches:
                    raise ValueError(
                        f"Duplicate carrier selection row {locality.member}:{code}"
                    )
                if _positive_decimal(row[5]) and _positive_decimal(row[6]):
                    matches.add(code)
    return matches


def _select_pfs_reference_codes(source_paths: Mapping[str, Path]) -> tuple[str, ...]:
    """Select a reproducible, non-hand-picked code sample from pinned sources."""

    eligible_sets = [
        _eligible_pfs_codes(source_paths[artifact_id])
        for artifact_id in _PFS_RVU_ARTIFACT_IDS
    ]
    candidates = set.intersection(*eligible_sets)

    carrier_artifact_ids = sorted({period[2] for period in _PFS_PERIODS})
    for artifact_id in carrier_artifact_ids:
        for locality in _PFS_LOCALITIES:
            candidates &= _carrier_reference_codes(
                source_paths[artifact_id], locality, candidates
            )

    selected: list[str] = []
    for lower, upper in _PFS_CODE_BANDS:
        band = [code for code in candidates if lower <= int(code) <= upper]
        if not band:
            raise ValueError(
                f"No eligible carrier-backed PFS code in numeric band {lower}-{upper}"
            )
        selected.append(
            min(
                band,
                key=lambda code: (
                    hashlib.sha256(
                        f"{_PFS_SELECTION_SALT}:{code}".encode("ascii")
                    ).digest(),
                    code,
                ),
            )
        )

    if len(set(selected)) != len(_PFS_CODE_BANDS):
        raise AssertionError("PFS source-derived selector returned duplicate codes")
    return tuple(selected)


def _carrier_amounts(
    path: Path,
    member: str,
    desired: set[tuple[str, str, str]],
) -> dict[tuple[str, str, str], tuple[Decimal, Decimal, int]]:
    matches: dict[tuple[str, str, str], tuple[Decimal, Decimal, int]] = {}
    with zipfile.ZipFile(path) as archive:
        if member not in archive.namelist():
            raise ValueError(f"Carrier member {member!r} is absent from {path.name}")
        with archive.open(member) as raw:
            rows = csv.reader(io.TextIOWrapper(raw, encoding="ascii"))
            for row_number, row in enumerate(rows, start=1):
                if len(row) < 7 or _text(row[0]) != "2024" or _text(row[4]):
                    continue
                key = (_text(row[1]), _text(row[2]), _text(row[3]).upper())
                if key not in desired:
                    continue
                if key in matches:
                    raise ValueError(f"Duplicate carrier reference row {member}:{key}")
                matches[key] = (
                    _decimal(row[5], label=f"{member} nonfacility amount"),
                    _decimal(row[6], label=f"{member} facility amount"),
                    row_number,
                )
    missing = desired - set(matches)
    if missing:
        raise ValueError(f"Missing carrier rows in {member}: {sorted(missing)}")
    return matches


def build_pfs_carrier_reference_cases(
    *,
    raw_dir: Path = RULEBOOK_RAW_DIR,
    manifest_path: Path = RULEBOOK_MANIFEST_PATH,
) -> tuple[PFSCarrierReferenceCase, ...]:
    """Build 40 expected PFS amounts from CMS's independent carrier files.

    Every one of the five supported snapshots contributes four source-selected
    codes across a rotating set of five localities, with both facility and
    nonfacility cases. No code or expected payment amount is hand-picked.
    """

    carrier_artifact_ids = tuple(sorted({period[2] for period in _PFS_PERIODS}))
    source_paths = _pinned_paths(
        (*_PFS_RVU_ARTIFACT_IDS, *carrier_artifact_ids),
        raw_dir=raw_dir,
        manifest_path=manifest_path,
    )
    selected_codes = _select_pfs_reference_codes(source_paths)

    selections: list[
        tuple[str, date, str, str, _CarrierLocality]
    ] = []
    for period_index, (version, service_date, artifact_id) in enumerate(_PFS_PERIODS):
        for code_index, code in enumerate(selected_codes):
            locality = _PFS_LOCALITIES[(code_index + period_index) % len(_PFS_LOCALITIES)]
            selections.append((version, service_date, artifact_id, code, locality))

    desired_by_member: dict[
        tuple[str, str], set[tuple[str, str, str]]
    ] = defaultdict(set)
    for _, _, artifact_id, code, locality in selections:
        desired_by_member[(artifact_id, locality.member)].add(
            (locality.carrier, locality.carrier_locality, code)
        )

    source_rows: dict[
        tuple[str, str],
        dict[tuple[str, str, str], tuple[Decimal, Decimal, int]],
    ] = {}
    for (artifact_id, member), desired in desired_by_member.items():
        source_rows[(artifact_id, member)] = _carrier_amounts(
            source_paths[artifact_id], member, desired
        )

    cases: list[PFSCarrierReferenceCase] = []
    for version, service_date, artifact_id, code, locality in selections:
        key = (locality.carrier, locality.carrier_locality, code)
        nonfacility, facility, row_number = source_rows[
            (artifact_id, locality.member)
        ][key]
        for setting, amount in (
            ("nonfacility", nonfacility),
            ("facility", facility),
        ):
            cases.append(
                PFSCarrierReferenceCase(
                    case_id=(
                        f"pfs-{version}-{code}-{locality.canonical_locality}-{setting}"
                    ),
                    code=code,
                    service_date=service_date,
                    rule_version=version,
                    locality=locality.canonical_locality,
                    setting=setting,
                    expected_amount=amount,
                    reference_artifact_id=artifact_id,
                    source_member=locality.member,
                    source_locator=(
                        f"{locality.member} row {row_number}; carrier "
                        f"{locality.carrier}, locality {locality.carrier_locality}, "
                        f"HCPCS {code}, blank modifier"
                    ),
                )
            )

    expected_count = len(_PFS_PERIODS) * len(selected_codes) * 2
    if len(cases) != expected_count or len({case.case_id for case in cases}) != len(cases):
        raise AssertionError(
            f"PFS reference-case design must yield {expected_count} unique cases"
        )
    return tuple(cases)


def _one_matching_name(
    names: Iterable[str],
    *,
    suffix: str,
    contains: str | None = None,
) -> str:
    matches = [
        name
        for name in names
        if name.lower().endswith(suffix.lower())
        and (contains is None or contains.lower() in name.lower())
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one member ending {suffix!r} containing {contains!r}; found {matches}"
        )
    return matches[0]


def _direct_workbook(path: Path) -> tuple[object, str]:
    with zipfile.ZipFile(path) as archive:
        member = _one_matching_name(archive.namelist(), suffix=".xlsx")
        workbook = load_workbook(
            io.BytesIO(archive.read(member)), read_only=True, data_only=True
        )
    return workbook, member


def _nested_ifc_workbook(path: Path) -> tuple[object, str]:
    with zipfile.ZipFile(path) as outer:
        nested_member = _one_matching_name(
            outer.namelist(), suffix=".zip", contains="IFC"
        )
        with zipfile.ZipFile(io.BytesIO(outer.read(nested_member))) as nested:
            workbook_member = _one_matching_name(
                nested.namelist(), suffix=".xlsx", contains="IFC"
            )
            workbook = load_workbook(
                io.BytesIO(nested.read(workbook_member)),
                read_only=True,
                data_only=True,
            )
    return workbook, f"{nested_member} > {workbook_member}"


def _normal_code(value: object, width: int) -> str:
    raw = _text(value)
    if raw.endswith(".0") and raw[:-2].isdigit():
        raw = raw[:-2]
    return raw.zfill(width) if raw.isdigit() else raw.upper()


def _rows_for_keys(
    sheet: object,
    keys: set[str],
    *,
    width: int,
) -> dict[str, tuple[int, tuple[object, ...]]]:
    found: dict[str, tuple[int, tuple[object, ...]]] = {}
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=3, values_only=True), start=3
    ):
        key = _normal_code(values[0] if values else None, width)
        if key not in keys:
            continue
        if key in found:
            raise ValueError(f"Duplicate source row for {key}")
        found[key] = (row_number, values)
    missing = keys - set(found)
    if missing:
        raise ValueError(f"Missing source rows: {sorted(missing)}")
    return found


def build_ipps_reference_cases(
    *,
    raw_dir: Path = RULEBOOK_RAW_DIR,
    manifest_path: Path = RULEBOOK_MANIFEST_PATH,
) -> tuple[IPPSReferenceCase, ...]:
    """Build eight IPPS reference cases directly from Tables 1, 2, and 5.

    The matrix crosses two MS-DRGs with high- and low-wage mainland providers
    in FY2024 and FY2025.  It therefore exercises both the 67.6 and 62 percent
    labor-share branches without using canonical rulebook parameters as inputs.
    """

    drgs = {"039", "470"}
    providers = {"010001", "050008"}
    artifact_ids = {
        f"cms_ipps_fy{fiscal_year}_{suffix}"
        for fiscal_year in (2024, 2025)
        for suffix in ("table1", "wage_tables", "table5")
    }
    paths = _pinned_paths(
        artifact_ids,
        raw_dir=raw_dir,
        manifest_path=manifest_path,
    )

    cases: list[IPPSReferenceCase] = []
    for fiscal_year, discharge_date in (
        (2024, date(2024, 9, 30)),
        (2025, date(2024, 10, 1)),
    ):
        prefix = f"cms_ipps_fy{fiscal_year}"
        table1_id = f"{prefix}_table1"
        wage_id = f"{prefix}_wage_tables"
        table5_id = f"{prefix}_table5"

        table1_workbook, table1_member = _direct_workbook(paths[table1_id])
        table5_workbook, table5_member = _direct_workbook(paths[table5_id])
        if fiscal_year == 2024:
            wage_workbook, wage_member = _direct_workbook(paths[wage_id])
            table1_sheet_name = "FY 2024 FINAL Table 1A-1E"
            table5_sheet_name = "FY 2024 Table 5 FR"
            wage_sheet_name = "Table 2 CN"
            wage_column_index = 6
            wage_column_letter = "G"
        else:
            wage_workbook, wage_member = _nested_ifc_workbook(paths[wage_id])
            table1_sheet_name = "FY 2025 IFC Table 1A-1E"
            table5_sheet_name = "FY 2025 Table 5 CN"
            wage_sheet_name = "Table 2 IFC"
            wage_column_index = 5
            wage_column_letter = "F"

        try:
            table1_sheet = table1_workbook[table1_sheet_name]
            table5_sheet = table5_workbook[table5_sheet_name]
            wage_sheet = wage_workbook[wage_sheet_name]
            drg_rows = _rows_for_keys(table5_sheet, drgs, width=3)
            wage_rows = _rows_for_keys(wage_sheet, providers, width=6)

            provider_values: dict[
                str,
                tuple[
                    Decimal,
                    Decimal,
                    Decimal,
                    Decimal,
                    str,
                    tuple[str, ...],
                ],
            ] = {}
            for provider_ccn in sorted(providers):
                wage_row_number, wage_values = wage_rows[provider_ccn]
                wage_index = _decimal(
                    wage_values[wage_column_index],
                    label=f"FY{fiscal_year} wage index {provider_ccn}",
                )
                transition = (
                    _optional_decimal(wage_values[6]) if fiscal_year == 2025 else None
                )
                out_migration = _optional_decimal(wage_values[16])
                if transition not in (None, Decimal("0")):
                    raise ValueError(
                        f"Reference provider {provider_ccn} has a transition factor"
                    )
                if out_migration not in (None, Decimal("0")):
                    raise ValueError(
                        f"Reference provider {provider_ccn} has out-migration adjustment"
                    )

                labor_share = Decimal("62") if wage_index <= 1 else Decimal("67.6")
                amount_row = 10 if labor_share == Decimal("62") else 5
                labor_cell = table1_sheet.cell(row=amount_row, column=1)
                nonlabor_cell = table1_sheet.cell(row=amount_row, column=2)
                labor_amount = _decimal(
                    labor_cell.value,
                    label=f"FY{fiscal_year} labor standardized amount",
                )
                nonlabor_amount = _decimal(
                    nonlabor_cell.value,
                    label=f"FY{fiscal_year} nonlabor standardized amount",
                )
                geography = _text(wage_values[11])
                if not geography:
                    raise ValueError(f"Reference provider {provider_ccn} lacks geography")
                provider_values[provider_ccn] = (
                    wage_index,
                    labor_share,
                    labor_amount,
                    nonlabor_amount,
                    geography,
                    (
                        f"{wage_member} > {wage_sheet_name}!"
                        f"{wage_column_letter}{wage_row_number}",
                        f"{table1_member} > {table1_sheet_name}!{labor_cell.coordinate}",
                        f"{table1_member} > {table1_sheet_name}!{nonlabor_cell.coordinate}",
                    ),
                )

            for ms_drg in sorted(drgs):
                drg_row_number, drg_values = drg_rows[ms_drg]
                relative_weight = _decimal(
                    drg_values[7], label=f"FY{fiscal_year} MS-DRG {ms_drg} weight"
                )
                for provider_ccn in sorted(providers):
                    (
                        wage_index,
                        labor_share,
                        labor_amount,
                        nonlabor_amount,
                        geography,
                        parameter_locators,
                    ) = provider_values[provider_ccn]
                    unrounded = (
                        labor_amount * wage_index + nonlabor_amount
                    ) * relative_weight
                    nine_places = unrounded.quantize(
                        _NINE_PLACES, rounding=ROUND_HALF_UP
                    )
                    expected = nine_places.quantize(_CENTS, rounding=ROUND_HALF_UP)
                    cases.append(
                        IPPSReferenceCase(
                            case_id=(
                                f"ipps-fy{fiscal_year}-{ms_drg}-{provider_ccn}"
                            ),
                            fiscal_year=f"FY{fiscal_year}",
                            ms_drg=ms_drg,
                            provider_ccn=provider_ccn,
                            geography=geography,
                            discharge_date=discharge_date,
                            quality_submitted=True,
                            meaningful_ehr_user=True,
                            relative_weight=relative_weight,
                            wage_index=wage_index,
                            labor_share=labor_share,
                            labor_amount=labor_amount,
                            nonlabor_amount=nonlabor_amount,
                            unrounded_amount=unrounded,
                            amount_nine_places=nine_places,
                            expected_amount=expected,
                            source_artifact_ids=(table1_id, wage_id, table5_id),
                            source_locators=(
                                f"{table5_member} > {table5_sheet_name}!H{drg_row_number}",
                                *parameter_locators,
                            ),
                        )
                    )
        finally:
            table1_workbook.close()
            table5_workbook.close()
            wage_workbook.close()

    if len(cases) != 8 or len({case.case_id for case in cases}) != len(cases):
        raise AssertionError("IPPS reference-case design must yield eight unique cases")
    return tuple(cases)


def _pfs_validation_store(
    store: RulebookStore,
    cases: Sequence[PFSCarrierReferenceCase],
) -> RulebookStore:
    """Retain every relevant interval while avoiding repeated full-table scans."""

    codes = {case.code for case in cases}
    localities = {case.locality for case in cases}
    assignments = store.code_assignments
    assignment_mask = (
        assignments["payment_system"].astype("string").str.upper().eq("PFS")
        & assignments["code"].astype("string").str.upper().isin(codes)
    )

    parameters = store.rule_parameters
    parameter_names = parameters["parameter_name"].astype("string")
    parameter_keys = parameters["key_value"].astype("string").str.upper()
    parameter_mask = parameters["rule_id"].astype("string").str.startswith(
        "pfs."
    ) & (
        parameter_names.eq("conversion_factor")
        | (
            parameter_names.isin({"work_gpci", "pe_gpci", "mp_gpci"})
            & parameter_keys.isin(localities)
        )
    )

    rules = store.payment_rules
    rule_mask = rules["rule_id"].astype("string").str.startswith("pfs.")
    return RulebookStore(
        payment_rules=rules.loc[rule_mask].copy(),
        rule_parameters=parameters.loc[parameter_mask].copy(),
        code_assignments=assignments.loc[assignment_mask].copy(),
        source_artifacts=store.source_artifacts,
    )


def _ipps_validation_store(
    store: RulebookStore,
    cases: Sequence[IPPSReferenceCase],
) -> RulebookStore:
    """Build an equivalent in-memory view containing only reference lookup keys."""

    drgs = {case.ms_drg for case in cases}
    providers = {case.provider_ccn for case in cases}
    assignments = store.code_assignments
    assignment_mask = (
        assignments["assignment_type"].astype("string").eq("ms_drg_weight")
        & assignments["code"].astype("string").str.zfill(3).isin(drgs)
    )

    parameters = store.rule_parameters
    parameter_names = parameters["parameter_name"].astype("string")
    provider_ids = parameters["provider_id"].astype("string").str.zfill(6)
    parameter_keys = parameters["key_value"].astype("string").str.zfill(6)
    standardized_mask = parameter_names.isin(
        {"standardized_labor_amount", "standardized_nonlabor_amount"}
    )
    provider_mask = parameter_names.isin(
        {
            "wage_index",
            "transitional_exception_wage_factor",
            "out_migration_adjustment",
        }
    ) & (provider_ids.isin(providers) | parameter_keys.isin(providers))

    parameter_mask = parameters["rule_id"].astype("string").str.startswith(
        "ipps."
    ) & (standardized_mask | provider_mask)

    rules = store.payment_rules
    rule_mask = rules["rule_id"].astype("string").str.startswith("ipps.")
    return RulebookStore(
        payment_rules=rules.loc[rule_mask].copy(),
        rule_parameters=parameters.loc[parameter_mask].copy(),
        code_assignments=assignments.loc[assignment_mask].copy(),
        source_artifacts=store.source_artifacts,
    )


def validate_pfs_carrier_references(
    store: RulebookStore,
    *,
    cases: Sequence[PFSCarrierReferenceCase] | None = None,
    raw_dir: Path = RULEBOOK_RAW_DIR,
    manifest_path: Path = RULEBOOK_MANIFEST_PATH,
) -> ReferenceValidationSummary:
    selected_cases = tuple(cases) if cases is not None else build_pfs_carrier_reference_cases(
        raw_dir=raw_dir, manifest_path=manifest_path
    )
    validation_store = _pfs_validation_store(store, selected_cases)
    carrier_artifact_ids = {
        case.reference_artifact_id for case in selected_cases
    }
    for case in selected_cases:
        trace = calculate_pfs(
            case.code,
            case.service_date,
            case.locality,
            case.setting,
            validation_store,
        )
        if trace.calculation_status is not CalculationStatus.CALCULATED:
            raise AssertionError(
                f"{case.case_id} did not calculate: {trace.to_dict()}"
            )
        calculation_source_ids = {
            _text(record.get("source_artifact_id"))
            for record in (
                *trace.selected_rule_versions,
                *trace.selected_parameters,
            )
            if _text(record.get("source_artifact_id"))
        }
        carrier_input_overlap = calculation_source_ids & carrier_artifact_ids
        if carrier_input_overlap:
            raise AssertionError(
                f"{case.case_id}: carrier artifact used as a calculation input: "
                f"{sorted(carrier_input_overlap)}"
            )
        if trace.calculated_amount != case.expected_amount:
            raise AssertionError(
                f"{case.case_id}: engine {trace.calculated_amount} != "
                f"carrier {case.expected_amount} at {case.source_locator}"
            )
        if not any(
            rule.get("rule_version") == case.rule_version
            for rule in trace.selected_rule_versions
        ):
            raise AssertionError(f"{case.case_id} selected the wrong rule version")
        validate_trace(trace)

    return ReferenceValidationSummary(
        payment_system="PFS",
        matched_cases=len(selected_cases),
        codes=tuple(sorted({case.code for case in selected_cases})),
        geographies=tuple(sorted({case.locality for case in selected_cases})),
        periods=tuple(sorted({case.rule_version for case in selected_cases})),
        settings=tuple(sorted({case.setting for case in selected_cases})),
        source_artifact_ids=tuple(
            sorted({case.reference_artifact_id for case in selected_cases})
        ),
    )


def validate_ipps_references(
    store: RulebookStore,
    *,
    cases: Sequence[IPPSReferenceCase] | None = None,
    raw_dir: Path = RULEBOOK_RAW_DIR,
    manifest_path: Path = RULEBOOK_MANIFEST_PATH,
) -> ReferenceValidationSummary:
    selected_cases = tuple(cases) if cases is not None else build_ipps_reference_cases(
        raw_dir=raw_dir, manifest_path=manifest_path
    )
    validation_store = _ipps_validation_store(store, selected_cases)
    for case in selected_cases:
        trace = calculate_ipps_base_payment(
            validation_store,
            ms_drg=case.ms_drg,
            discharge_date=case.discharge_date,
            provider_ccn=case.provider_ccn,
            quality_submitted=case.quality_submitted,
            meaningful_ehr_user=case.meaningful_ehr_user,
        )
        if trace.calculation_status is not CalculationStatus.CALCULATED:
            raise AssertionError(
                f"{case.case_id} did not calculate: {trace.to_dict()}"
            )
        expected_components: Mapping[str, Decimal] = {
            "ms_drg_relative_weight": case.relative_weight,
            "wage_index": case.wage_index,
            "labor_share_percent": case.labor_share,
            "labor_standardized_amount": case.labor_amount,
            "nonlabor_standardized_amount": case.nonlabor_amount,
            "unrounded_base_operating_payment": case.unrounded_amount,
            "base_operating_payment_9dp": case.amount_nine_places,
        }
        for component, expected in expected_components.items():
            if trace.components.get(component) != expected:
                raise AssertionError(
                    f"{case.case_id} {component}: "
                    f"{trace.components.get(component)} != {expected}"
                )
        if trace.calculated_amount != case.expected_amount:
            raise AssertionError(
                f"{case.case_id}: engine {trace.calculated_amount} != "
                f"raw-table reference {case.expected_amount}"
            )
        validate_trace(trace)

    return ReferenceValidationSummary(
        payment_system="IPPS",
        matched_cases=len(selected_cases),
        codes=tuple(sorted({case.ms_drg for case in selected_cases})),
        geographies=tuple(sorted({case.geography for case in selected_cases})),
        periods=tuple(sorted({case.fiscal_year for case in selected_cases})),
        labor_share_branches=tuple(
            sorted({format(case.labor_share, "f") for case in selected_cases})
        ),
        source_artifact_ids=tuple(
            sorted(
                {
                    artifact_id
                    for case in selected_cases
                    for artifact_id in case.source_artifact_ids
                }
            )
        ),
    )


def validate_expanded_references(
    store: RulebookStore,
) -> tuple[ReferenceValidationSummary, ReferenceValidationSummary]:
    """Run both expanded reference matrices for integration by the main validator."""

    return (
        validate_pfs_carrier_references(store),
        validate_ipps_references(store),
    )


__all__ = [
    "IPPSReferenceCase",
    "PFSCarrierReferenceCase",
    "ReferenceValidationSummary",
    "build_ipps_reference_cases",
    "build_pfs_carrier_reference_cases",
    "validate_expanded_references",
    "validate_ipps_references",
    "validate_pfs_carrier_references",
]
